"""Connected synthetic Portal journey; invoked by portal_local_audit.py only."""
from __future__ import annotations

from contextlib import ExitStack
import csv
from datetime import timedelta
from hashlib import sha256
import hmac
from io import BytesIO, StringIO
import json
import time
import traceback
from urllib.parse import urlencode
from unittest.mock import patch


def run(output, *, source='sysup', intake_count=2, completion_count=None):
    from django.conf import settings
    from django.contrib.auth import get_user_model
    from django.core.files.base import ContentFile
    from django.core.files.uploadedfile import SimpleUploadedFile
    from django.test import Client
    from django.urls import reverse
    from django.utils import timezone
    from openpyxl import Workbook, load_workbook
    from reportlab.pdfgen.canvas import Canvas
    from core.models import (
        AccessGrant, GroupSheetConfiguration, IntegrationOperation,
        JawabuFarmerMaster, JawabuFarmerUploadBatch, PaymentDocumentTemplate,
        RequisitionBatch, RequisitionTemplate, UserProfile,
    )
    from core.tests import FakeMasterDataSheet, FakeJawabuService
    from core.tests_payment_documents import synthetic_payment_template_bytes
    from core.services.order_approval import GoogleDriveMediaStorage
    from core.services.portal_imports import farmup_revision_token, attempt_import_archive
    from core.services.portal_publication import queued_publication_operations, attempt_publication
    from core.services.jawabu_pipeline import current_pipeline_state_label, JBL_FORWARD_STATUS
    from hb_operations.models import HomeBiogasAction
    from payments.models import PaymentBatch
    from decimal import Decimal
    from core.services.requisition_partners import fulfillment_partner_for_farmer

    completion_count = completion_count or intake_count
    varied = intake_count > 2 or source != 'sysup'
    started = time.monotonic()

    evidence = {'steps': [], 'status': 'running', 'network': 'blocked',
                'business_logic_patched': False, 'database': settings.DATABASES['default']['NAME'],
                'source': source, 'intake_count': intake_count, 'completion_count': completion_count}
    def record(name, **details):
        evidence['steps'].append({'step': name, **details})
        evidence['elapsed_seconds'] = round(time.monotonic() - started, 2)
        (output / 'journey.json').write_text(json.dumps(evidence, indent=2, default=str), encoding='utf-8')
        print(name, json.dumps(details, default=str), flush=True)

    assert not JawabuFarmerMaster.objects.exists(), 'Journey requires its dedicated empty database'
    settings.TELEGRAM_BOT_TOKEN = '000000:synthetic-local-audit-only'
    settings.GOOGLE_DRIVE_MEDIA_FOLDER_ID = 'local-audit-root'
    settings.PORTAL_PUBLICATION_MIN_SPACING_SECONDS = 0
    settings.RUNNING_TESTS = True
    drive_dir = output / 'drive'
    sheet_dir = output / 'sheets'
    drive_dir.mkdir(exist_ok=True)
    sheet_dir.mkdir(exist_ok=True)
    drive_files = []
    def upload(storage, data, *, filename, mime_type, **kwargs):
        if hasattr(data, 'read'):
            data.seek(0)
            data = data.read()
        data = bytes(data)
        file_id = sha256(data).hexdigest()
        (drive_dir / file_id).write_bytes(data)
        drive_files.append({'id': file_id, 'filename': filename, 'mime_type': mime_type, 'size': len(data)})
        (drive_dir / 'manifest.json').write_text(json.dumps(drive_files, indent=2), encoding='utf-8')
        return file_id, f'https://drive.invalid/local/{file_id}'
    def download(storage, file_id):
        return (drive_dir / file_id).read_bytes()

    headers = ['No.', 'Customer Name', 'National ID', 'Primary Phone', 'County', 'Sub-County',
               'Village', 'Branch', 'Case ID', 'Current Pipeline State', 'JBL Visit Date',
               'Credit Analysis', 'Final Decision', 'Order No', 'Invoice Number', 'Payment No',
               'Installation Status', 'Installation Date', 'Commissioning Status', 'Commissioning Date',
               'Balance Due', 'Deposit Paid to HBG', 'CUSTOMER NO', 'NAME (IMAB)', 'Payment Product',
               'Mode of Payment', 'Requisition Date']
    sheets = {}
    for tab in ('Master Data', 'Eco-conserve', 'Orders'):
        sheet = FakeMasterDataSheet(headers)
        sheet.values = [list(headers)]
        sheets[tab] = sheet
    def sheet_service(*, sheet_id, sheet_name, **kwargs):
        assert sheet_id == 'local-audit-sheet', 'Unexpected spreadsheet destination'
        return FakeJawabuService(sheets[sheet_name])
    def save_sheets():
        for tab, sheet in sheets.items():
            with (sheet_dir / f'{tab}.csv').open('w', newline='', encoding='utf-8') as stream:
                csv.writer(stream).writerows(sheet.values)
    def drain(name):
        attempts = 0
        while (operation := queued_publication_operations().first()) is not None:
            result = attempt_publication(operation)
            attempts += 1
            assert not result.get('error'), result
            assert attempts <= intake_count * 30 + 100, 'Publication queue stopped progressing'
        save_sheets()
        record(name, attempts=attempts, rows={tab: len(sheet.values) - 1 for tab, sheet in sheets.items()})

    group = GroupSheetConfiguration.objects.create(
        group_id='local-synthetic-portal', display_name='Synthetic Portal audit', enabled=True,
        sheet_id='local-audit-sheet', sheet_name='Master Data',
        workflow={'type': 'jawabu_homebiogas', 'master_sync_enabled': True,
                  'master_sheet_id': 'local-audit-sheet', 'master_sheet_name': 'Master Data',
                  'eco_conserve_sheet_name': 'Eco-conserve', 'master_header_row': 1,
                  'master_data_start_row': 2},
    )
    clients = {}
    for number, role in enumerate(('IT', 'OPERATIONS_ADMIN', 'JBL_OFFICER', 'CREDIT_ANALYST', 'BUSINESS_ADMIN', 'HB_STAFF'), 1):
        actor = get_user_model().objects.create_user(username=f'synthetic-{role.lower()}', first_name=f'Synthetic {role}')
        UserProfile.objects.update_or_create(user=actor, defaults={'telegram_id': str(900000000 + number)})
        AccessGrant.objects.create(user=actor, workflow='jawabu_portal', role=role, group_configuration=group)
        fields = {'auth_date': str(int(time.time())), 'user': json.dumps({'id': 900000000 + number, 'first_name': actor.first_name})}
        secret = hmac.new(b'WebAppData', settings.TELEGRAM_BOT_TOKEN.encode(), sha256).digest()
        fields['hash'] = hmac.new(secret, '\n'.join(f'{key}={value}' for key, value in sorted(fields.items())).encode(), sha256).hexdigest()
        clients[role] = Client(HTTP_X_TELEGRAM_INIT_DATA=urlencode(fields))

    def api(role, name, data=None, *, args=(), expected=200, multipart=False, key=None):
        url = reverse(name, args=args)
        request_started = time.monotonic()
        if data is None:
            response = clients[role].get(url)
        else:
            payload = {**data, 'client_request_id': key or f'local-audit-{len(evidence["steps"]):03}-{name}'}
            headers = {'HTTP_X_REQUEST_ID': payload['client_request_id']}
            response = clients[role].post(url, payload, **headers) if multipart else clients[role].post(url, json.dumps(payload), content_type='application/json', **headers)
        try:
            result = response.json()
        except ValueError:
            result = {'body': response.content[:300].decode(errors='replace')}
        record(name, role=role, status=response.status_code, expected=expected,
               ok=result.get('ok'), error=result.get('error'),
               duration_ms=round((time.monotonic() - request_started) * 1000, 2))
        assert response.status_code == expected, result
        if expected < 300:
            assert result.get('ok', True), result
        return result

    def csv_bytes(columns, rows):
        stream = StringIO()
        writer = csv.writer(stream)
        writer.writerow(columns)
        writer.writerows(rows)
        return stream.getvalue().encode()

    def pdf(text):
        stream = BytesIO()
        canvas = Canvas(stream)
        canvas.setTitle('Synthetic local Portal audit')
        for index, line in enumerate(text.splitlines()):
            canvas.drawString(45, 800 - index * 18, line)
        canvas.save()
        return stream.getvalue()

    def photo(number):
        from PIL import Image, ImageDraw
        image = Image.new('RGB', (800, 600), (30 * number, 60, 150))
        ImageDraw.Draw(image).text((50, 50), f'SYNTHETIC CAPTURE {number}', fill='white')
        stream = BytesIO()
        image.save(stream, format='JPEG')
        return SimpleUploadedFile(f'synthetic-{number}.jpg', stream.getvalue(), content_type='image/jpeg')

    try:
        with ExitStack() as stack:
            stack.enter_context(patch.object(GoogleDriveMediaStorage, 'upload', upload))
            stack.enter_context(patch.object(GoogleDriveMediaStorage, 'download', download))
            stack.enter_context(patch('core.services.sheets.GoogleSheetsService.get_instance', side_effect=sheet_service))
            # Telegram transport is a local sink; authentication remains real HMAC verification.
            stack.enter_context(patch('requests.post', side_effect=OSError('Synthetic audit: Telegram delivery disabled')))
            api('JBL_OFFICER', 'portal_meta')
            from requisitions.models import OrderSequenceState
            from payments.models import PaymentSequenceState
            for partner in ('HB', 'ECOCONSERVE'):
                OrderSequenceState.objects.create(group_configuration=group, partner=partner, next_number=1)
                workbook = Workbook()
                ws = workbook.active
                ws.append(['Synthetic requisition'])
                ws['A4'], ws['H4'] = 'Date:', 'Order No:'
                for col, title in enumerate(('No.', 'Name of the Customer', 'Contact No', 'ID No', 'Credit Analysis', 'Callup Comment', 'County', 'Location', 'HBG', 'JBL', 'Sales Person'), 1):
                    ws.cell(7, col, title)
                for row in range(9, 14):
                    ws.cell(row, 1, row - 8)
                ws.cell(14, 1, 'Total')
                stream = BytesIO()
                workbook.save(stream)
                template = RequisitionTemplate.objects.create(name=f'Synthetic {partner}', fulfillment_partner=partner, is_active=True)
                template.file.save(f'synthetic-{partner}.xlsx', ContentFile(stream.getvalue()))
            PaymentSequenceState.objects.create(group_configuration=group, next_number=1)
            template = PaymentDocumentTemplate.objects.create(name='Synthetic payment', is_active=True)
            template.file.save('synthetic-payment.xlsx', ContentFile(synthetic_payment_template_bytes()))

            identities = [('SYNTHETIC HB APPLICANT', '99000001', '254700000001', 'Kiambu', 'Limuru', 'Synthetic HB'),
                          ('SYNTHETIC ECO APPLICANT', '99000002', '254700000002', 'Nakuru', 'Nakuru', 'Eco-conserve')]
            if varied:
                name_variants = ('APPLICANT', "O'NEIL", 'MWÄNGI', 'KAMAU, TEST', 'A VERY LONG ' + 'NAME ' * 24)
                identities = [(f'SYNTHETIC {name_variants[i % len(name_variants)]} {i:05}',
                               str(99000000 + i), f'2547{i:08}',
                               'Kiambu' if i % 2 else 'Nakuru', 'Limuru' if i % 2 else 'Nakuru',
                               'Synthetic HB' if i % 2 else 'Eco-conserve') for i in range(1, intake_count + 1)]

            def phone_format(value, index):
                return (value, '+' + value, '0' + value[3:], '00254' + value[3:],
                        '+254 (0) ' + value[3:6] + '-' + value[6:9] + '-' + value[9:],
                        '005' + value[3:])[index % 6] if varied else value

            from datetime import date
            date_cases = (date(2024, 2, 29), date(2025, 12, 31), date(2026, 1, 1),
                          date(2026, 4, 30), date(2026, 9, 30), timezone.localdate())

            def visit_day(index):
                return date_cases[index % len(date_cases)] if varied else timezone.localdate()

            def captures():
                return {'client_id_front': photo(1), 'client_id_back': photo(2),
                        'laf_page_1': photo(3), 'laf_page_2': photo(4), 'jbl_visit_photo_files': photo(5)}

            deposits = ('', '0', '1', '4999.99', '5000', '5000.01', '10000', '100000.00')
            sys_batches = []
            if source == 'sysup':
                sysup = csv_bytes(['Customer ID', 'Name', 'Mobile No', 'ID NO', 'Branch', 'Loan Officer', 'Product Name', 'LGF Balance'],
                                  [[str(990001 + i), name, phone_format(phone, i), identity, branch, 'Synthetic officer', 'HomeBiogas', str(i % 4 * 2500) if varied else '0']
                                   for i, (name, identity, phone, county, branch, sales) in enumerate(identities)])
                staged = api('IT', 'portal_import_stage', {'file': SimpleUploadedFile('synthetic-sysup.csv', sysup)}, args=('sysup',), multipart=True, expected=201)
                sys_batches.append(JawabuFarmerUploadBatch.objects.get(pk=staged['batch']['id']))
                assert not JawabuFarmerMaster.objects.exists()
                record('sysup_does_not_create_leads', rows=intake_count, uploads=len(sys_batches))
            if source == 'office':
                for i, (name, identity, phone, county, branch, sales) in enumerate(identities):
                    direct_visit = i % 10 == 0
                    created = api('JBL_OFFICER', 'portal_create_and_complete_jbl_lead', {
                        'customer_name': name, 'national_id': identity, 'primary_phone': phone_format(phone, i),
                        'county': county, 'sub_county': 'Limuru' if county == 'Kiambu' else 'Nakuru Town East',
                        'village': 'Synthetic village', 'hb_sales_person': sales,
                        'deposit_paid_hbg': deposits[4 + i % 4] if direct_visit else deposits[i % 4],
                        **({'visit_date': visit_day(i).isoformat(),
                            'visit_status': JBL_FORWARD_STATUS,
                            'location_unavailable_reason': 'Synthetic audit without GPS', **captures()} if direct_visit else {}),
                    }, multipart=True, expected=200 if direct_visit else 201, key=f'local-office-{i}')
                    if direct_visit:
                        assert created['farmer']['jbl_visit_date']
                    else:
                        assert created['deposit_hold'] and not created['visit_logged']
                record('office_intake_without_sysup', rows=intake_count)
            else:
                farmup = csv_bytes(['Full Name', 'ID NUMBER', 'HBG Hub', 'Branch', 'Mobile', 'Phone', 'Actual Receipts', 'Sign Date', 'Sales Person'],
                                   [[name, identity, county, branch, phone_format(phone, i), '254711' + phone[-6:],
                                     deposits[i % len(deposits)] if varied else '5000',
                                     (visit_day(i) - timedelta(days=2 + i % 180)).strftime('%d/%m/%Y'), sales]
                                    for i, (name, identity, phone, county, branch, sales) in enumerate(identities)])
                staged = api('OPERATIONS_ADMIN', 'portal_farmup_stage', {'file': SimpleUploadedFile('synthetic-farmup.csv', farmup), 'period': timezone.localdate().strftime('%Y-%m')}, multipart=True, expected=201)
                farm_batch = JawabuFarmerUploadBatch.objects.get(pk=staged['batch']['id'])
                rows = [{'row_id': row['row_id'], 'approved': True, 'warning_acknowledged': True} for row in farm_batch.parsed_rows]
                payload = {'rows': rows, 'revision_token': farmup_revision_token(farm_batch)}
                validated = api('OPERATIONS_ADMIN', 'portal_farmup_validate', payload, args=(farm_batch.pk,))
                assert validated['counts']['selected'] == intake_count, validated['counts']
                api('OPERATIONS_ADMIN', 'portal_farmup_commit', payload, args=(farm_batch.pk,), key='local-farmup-commit')
                api('OPERATIONS_ADMIN', 'portal_farmup_commit', payload, args=(farm_batch.pk,), key='local-farmup-commit')
                for batch in [*sys_batches, farm_batch]:
                    operation = IntegrationOperation.objects.get(source_id=str(batch.pk), operation_type='portal_import_drive_archive')
                    archive_result = attempt_import_archive(str(operation.pk))
                    assert not archive_result.get('error'), archive_result
            assert JawabuFarmerMaster.objects.count() == intake_count
            if source == 'sysup':
                from core.services.system_export import resolve_system_export_row
                for sys_batch in sys_batches:
                    while sys_batch.parsed_rows:
                        sys_rows = []
                        for row in sys_batch.parsed_rows[:settings.SYSUP_PROCESSING_CHUNK_SIZE]:
                            resolved = dict(row) if row.get('Matched Farmer ID') else resolve_system_export_row(dict(row))
                            farmer = JawabuFarmerMaster.objects.get(national_id=row['ID NO'])
                            sys_rows.append({**resolved, 'approved': True, 'Matched Farmer ID': str(farmer.pk), 'case_revision': farmer.workflow_revision})
                        result = api('OPERATIONS_ADMIN', 'portal_import_commit', {'rows': sys_rows, 'portal_revision': sys_batch.portal_revision}, args=(sys_batch.pk,))
                        assert result['result']['committed'] > 0, result
                        sys_batch.refresh_from_db()
            else:
                assert not JawabuFarmerMaster.objects.exclude(customer_no='').exists()
                assert not JawabuFarmerUploadBatch.objects.filter(import_kind='system_export').exists()
                record('no_sysup_or_initial_system_identity', source=source, rows=intake_count)
            drain('initial_sheet_publication')

            orders, receipts = [], []
            for index, (name, identity, phone, county, branch, sales) in enumerate(identities[:completion_count]):
                farmer = JawabuFarmerMaster.objects.get(national_id=identity)
                if varied and (farmer.deposit_paid_hbg is None or farmer.deposit_paid_hbg < Decimal('5000')):
                    api('JBL_OFFICER', 'portal_complete_jbl_visit', {'workflow_revision': farmer.workflow_revision}, args=(farmer.pk,), multipart=True, expected=400)
                    api('HB_STAFF', 'portal_hb_deposit_update', {'amount': '5000', 'expected_revision': farmer.workflow_revision}, args=(farmer.pk,))
                    farmer.refresh_from_db()
                api('OPERATIONS_ADMIN', 'portal_set_final_decision', {'final_decision': 'Approved', 'workflow_revision': farmer.workflow_revision}, args=(farmer.pk,), expected=403)
                visit = {'workflow_revision': farmer.workflow_revision, 'visit_date': visit_day(index).isoformat(),
                         'visit_status': JBL_FORWARD_STATUS, 'officer': 'Synthetic officer', 'county': county,
                         'sub_county': 'Limuru' if county == 'Kiambu' else 'Nakuru Town East', 'village': 'Synthetic village',
                         'location_unavailable_reason': 'Synthetic local audit without device GPS',
                         'client_id_front': photo(1), 'client_id_back': photo(2),
                         'laf_page_1': photo(3), 'laf_page_2': photo(4),
                         'jbl_visit_photo_files': photo(5)}
                if not farmer.jbl_visit_date:
                    api('JBL_OFFICER', 'portal_complete_jbl_visit', visit, args=(farmer.pk,), multipart=True)
                farmer.refresh_from_db()
                api('CREDIT_ANALYST', 'portal_set_credit_decision', {'workflow_revision': farmer.workflow_revision - 1,
                    'decision': 'Approved', 'imab_created': 'Yes', 'customer_no': farmer.customer_no}, args=(farmer.pk,), expected=409)
                api('CREDIT_ANALYST', 'portal_set_credit_decision', {'workflow_revision': farmer.workflow_revision,
                    'decision': 'Approved', 'imab_created': 'Yes', 'customer_no': farmer.customer_no or str(880000 + index)}, args=(farmer.pk,))
                farmer.refresh_from_db()
                api('BUSINESS_ADMIN', 'portal_set_final_decision', {'workflow_revision': farmer.workflow_revision,
                    'final_decision': 'Approved', 'repayment_date': '10', 'repayment_tenor': '6', 'decision_comment': 'Synthetic order approval'}, args=(farmer.pk,))
                preview = api('OPERATIONS_ADMIN', 'portal_requisition_preview', {'farmer_ids': [str(farmer.pk)], 'requisition_date': timezone.localdate().isoformat()})
                generated = api('OPERATIONS_ADMIN', 'portal_requisition_finalize', {'preview_token': preview['preview_token']})
                order = RequisitionBatch.objects.get(pk=generated['batch']['id'])
                from core.services.invoice_parser import official_requisition_eligibility
                farmer.refresh_from_db()
                assert official_requisition_eligibility(farmer)['code'] == 'invoice_order_unsigned'
                record('unsigned_order_blocks_invoice', order=order.order_number)
                (output / f'{order.order_number}.xlsx').write_bytes(bytes(order.file_content))
                assert load_workbook(BytesIO(bytes(order.file_content))).active.cell(9, 2).value == farmer.customer_name
                scan = SimpleUploadedFile('synthetic-signed-order.pdf', pdf('SYNTHETIC signed and stamped order\n' + order.order_number), content_type='application/pdf')
                api('BUSINESS_ADMIN', 'portal_document_physical_signoff_upload', {'signed_scan': scan, 'attested_complete': 'true'}, args=('requisition', order.pk), multipart=True)
                assert HomeBiogasAction.objects.filter(farmer=farmer).exists()
                orders.append(order)
                balance = Decimal('44500.00') + (Decimal(index % 8) * Decimal('137.25') if varied else 0)
                gross = balance + Decimal('9500')
                invoice_pdf = pdf(f'HOMEBIOGAS VENTURES LIMITED\nBILL TO\n{farmer.customer_name}\n{identity}\n{phone_format(phone, index)}\nINVOICE {990001 + len(orders)}\nDATE {timezone.localdate():%d/%m/%Y}\nSUBTOTAL {gross:,.2f}\nDISCOUNT -4,500.00\nTOTAL {gross - Decimal("4500"):,.2f}\nPAYMENT 5,000.00\nBALANCE DUE {balance:,.2f}')
                uploaded = api('OPERATIONS_ADMIN', 'portal_invoice_pool_upload', {
                    'file': SimpleUploadedFile(f'synthetic-invoice-{identity}.pdf', invoice_pdf, content_type='application/pdf')}, multipart=True)
                assert uploaded['auto_matched_count'] == 1, uploaded
                receipts.append(uploaded['receipt_batch'])
                farmer.refresh_from_db()
                assert farmer.balance_due == balance, str(farmer.balance_due)
                record('invoice_parsed_and_matched', farmer_id=str(farmer.pk), invoice=farmer.invoice_number, balance=str(farmer.balance_due))

            farmers = list(JawabuFarmerMaster.objects.filter(national_id__in=[item[1] for item in identities[:completion_count]]).order_by('national_id'))
            payment_modes = {str(farmer.pk): 'LOAN-JAWABU' if i % 2 == 0 else 'CASH' for i, farmer in enumerate(farmers)}
            created = api('OPERATIONS_ADMIN', 'portal_invoice_receipt_batch_payment', {
                'revision': receipts[0]['revision'], 'payment_modes': {str(farmers[0].pk): 'LOAN-JAWABU'}}, args=(receipts[0]['id'],), expected=201)
            batch = PaymentBatch.objects.get(pk=created['batch']['id'])
            api('OPERATIONS_ADMIN', 'portal_payment_batch_cases', {'farmer_ids': [str(farmer.pk) for farmer in farmers],
                'payment_modes': payment_modes, 'revision': batch.revision}, args=(batch.pk,))
            batch.refresh_from_db()
            api('OPERATIONS_ADMIN', 'portal_payment_batch_submit', {'revision': batch.revision}, args=(batch.pk,))
            batch.refresh_from_db()
            for farmer in farmers:
                api('OPERATIONS_ADMIN', 'portal_payment_batch_case_review', {'decision': 'approved', 'comment': 'Forbidden synthetic attempt', 'revision': batch.revision}, args=(batch.pk, farmer.pk), expected=403)
                api('BUSINESS_ADMIN', 'portal_payment_batch_case_review', {'decision': 'approved', 'comment': 'Synthetic payment approval', 'revision': batch.revision}, args=(batch.pk, farmer.pk))
                batch.refresh_from_db()
            api('BUSINESS_ADMIN', 'portal_payment_batch_generate', {'revision': batch.revision}, args=(batch.pk,))
            batch.refresh_from_db()
            document = batch.current_document
            (output / 'payment.xlsx').write_bytes(bytes(document.file_content))
            from core.services.payment_documents import payment_template_layout
            workbook = load_workbook(BytesIO(bytes(document.file_content)))
            layout = payment_template_layout(workbook)
            from core.services.requisition import ordered_document_farmers
            expected_order = ordered_document_farmers(farmers, eco_first=True)
            self_names = [workbook[layout.sheet_name].cell(layout.data_start_row + i, layout.columns['name']).value
                          for i in range(len(farmers))]
            assert self_names == [farmer.customer_name for farmer in expected_order], self_names
            api('BUSINESS_ADMIN', 'portal_document_physical_signoff_upload', {'signed_scan': SimpleUploadedFile('synthetic-signed-payment.pdf', pdf('SYNTHETIC signed payment'), content_type='application/pdf'),
                'attested_complete': 'true'}, args=('payment', document.pk), multipart=True)
            batch.refresh_from_db()
            assert batch.status == PaymentBatch.STATUS_COMPLETED, batch.status
            record('payment_completed', number=batch.payment_number, cases=len(farmers))
            for i, farmer in enumerate(farmers):
                action = HomeBiogasAction.objects.get(farmer=farmer)
                installed_on = timezone.localdate() - timedelta(days=45 + i % 365) if varied else timezone.localdate()
                commissioned_on = installed_on + timedelta(days=7 if i % 2 else 30) if varied else timezone.localdate()
                api('HB_STAFF', 'portal_hb_action_transition', {'workstream': 'commissioning', 'commissioning_date': timezone.localdate().isoformat(), 'revision': action.revision}, args=(farmer.pk,), expected=400)
                api('HB_STAFF', 'portal_hb_action_transition', {'workstream': 'installation', 'installation_status': 'installed', 'installation_date': installed_on.isoformat(), 'revision': action.revision}, args=(farmer.pk,))
                action.refresh_from_db()
                api('HB_STAFF', 'portal_hb_action_transition', {'workstream': 'commissioning', 'commissioning_date': (installed_on - timedelta(days=1)).isoformat(), 'early_commissioning_acknowledged': True, 'revision': action.revision}, args=(farmer.pk,), expected=400)
                if (commissioned_on - installed_on).days < 21:
                    api('HB_STAFF', 'portal_hb_action_transition', {'workstream': 'commissioning', 'commissioning_date': commissioned_on.isoformat(), 'revision': action.revision}, args=(farmer.pk,), expected=400)
                api('HB_STAFF', 'portal_hb_action_transition', {'workstream': 'commissioning', 'commissioning_date': commissioned_on.isoformat(), 'early_commissioning_acknowledged': True, 'revision': action.revision}, args=(farmer.pk,))
                action.refresh_from_db()
                assert action.commissioning_status == 'commissioned'
                farmer.refresh_from_db()
                record('commissioned', farmer_id=str(farmer.pk), pipeline=current_pipeline_state_label(farmer))
            drain('final_sheet_publication')
            for tab in ('Master Data', 'Eco-conserve'):
                values = sheets[tab].values
                expected_farmers = [farmer for farmer in JawabuFarmerMaster.objects.all()
                                    if (fulfillment_partner_for_farmer(farmer) == 'ECOCONSERVE') == (tab == 'Eco-conserve')]
                assert len(values) - 1 == len(expected_farmers), (tab, len(values), len(expected_farmers))
                by_id = {row[values[0].index('Master Record ID')]: row for row in values[1:]}
                assert len(by_id) == len(expected_farmers)
                for farmer in farmers:
                    if (fulfillment_partner_for_farmer(farmer) == 'ECOCONSERVE') != (tab == 'Eco-conserve'):
                        continue
                    row = by_id[str(farmer.pk)]
                    assert row[values[0].index('Current Pipeline State')] == 'COMMISSIONED'
                    assert str(row[values[0].index('Payment No')]) == '1'
                    from decimal import ROUND_HALF_UP
                    assert Decimal(str(row[values[0].index('Balance Due')])) == farmer.balance_due.quantize(Decimal('1'), rounding=ROUND_HALF_UP)
            from datetime import datetime
            sheet_timestamp = sheets['Master Data'].values[1][sheets['Master Data'].values[0].index('Last Updated At')]
            displayed = datetime.strptime(sheet_timestamp, '%d-%b-%Y %H:%M').replace(tzinfo=timezone.get_current_timezone())
            record('sheet_timestamp_zone_probe', displayed=sheet_timestamp,
                   nairobi_now=timezone.localtime().strftime('%d-%b-%Y %H:%M'),
                   lag_minutes=round((timezone.now() - displayed).total_seconds() / 60))
            assert abs((timezone.now() - displayed).total_seconds()) < 120
            assert all(probe['passes_expected_behavior'] for probe in probe_invoice_phones(output))
            assert not IntegrationOperation.objects.exclude(status=IntegrationOperation.STATUS_SUCCEEDED).exists()
            evidence['status'] = 'passed'
            record('completed', customers=JawabuFarmerMaster.objects.count(), drive_uploads=len(drive_files),
                   publication_statuses=list(IntegrationOperation.objects.values_list('status', flat=True)))
            return 0
    except Exception as exc:
        evidence['status'] = 'failed'
        record('failure', error=str(exc), traceback=traceback.format_exc())
        save_sheets()
        return 1


def probe_invoice_phones(output):
    """Characterize parsing/matching without persisting any invoice or case."""
    from core.models import JawabuFarmerMaster, ParsedInvoice
    from core.services.invoice_parser import parse_invoice_text, _exact_invoice_identity_match
    farmer = JawabuFarmerMaster(customer_name='SYNTHETIC APPLICANT', national_id='99000009', primary_phone='254712345678')
    results = []
    for printed in ('+254712345678', '0712345678', '254712345678', '+254798765432', '254798765432'):
        parsed = parse_invoice_text(f'HOMEBIOGAS VENTURES LIMITED\nBILL TO\nSYNTHETIC APPLICANT\n99000009\n{printed}\nINVOICE 999900\nDATE 26/10/2026\nTOTAL 100.00\nBALANCE DUE 100.00', 1)
        assert parsed is not None
        invoice = ParsedInvoice(customer_name=parsed['customer_name'], customer_id=parsed['customer_id'], customer_phone=parsed['customer_phone'])
        matches = _exact_invoice_identity_match(invoice, farmer)
        expected = '98765432' not in printed
        results.append({'printed_phone': printed, 'parsed_phone': parsed['customer_phone'],
                        'matched': matches, 'expected_match': expected, 'passes_expected_behavior': matches == expected})
    (output / 'invoice-phone-probes.json').write_text(json.dumps(results, indent=2), encoding='utf-8')
    return results


def verify_retained_journey(output):
    """Independently reconcile retained DB facts, local Sheet cells and workbook bytes."""
    from decimal import Decimal, ROUND_HALF_UP
    from django.utils import timezone
    from openpyxl import load_workbook
    from core.models import IntegrationOperation, JawabuFarmerMaster, RequisitionBatch, ParsedInvoice
    from core.services.requisition_partners import fulfillment_partner_for_farmer
    from core.services.jawabu_case_reference import display_case_reference
    from hb_operations.models import HomeBiogasAction
    from payments.models import PaymentBatch

    evidence = json.loads((output / 'journey.json').read_text(encoding='utf-8'))
    expected = evidence['completion_count']
    assert expected >= 2
    farmers = list(JawabuFarmerMaster.objects.order_by('national_id'))
    assert len(farmers) == evidence['intake_count']
    completed = farmers[:expected]
    assert RequisitionBatch.objects.count() == expected
    assert ParsedInvoice.objects.filter(status='matched').count() == expected
    assert HomeBiogasAction.objects.filter(commissioning_status='commissioned').count() == expected
    batch = PaymentBatch.objects.get(status=PaymentBatch.STATUS_COMPLETED)
    assert batch.case_memberships.filter(is_active=True).count() == expected
    assert not IntegrationOperation.objects.exclude(status='succeeded').exists()
    assert not [step for step in evidence['steps'] if 'expected' in step and step['status'] != step['expected']]
    sheet_rows = 0
    for tab in ('Master Data', 'Eco-conserve'):
        with (output / 'sheets' / f'{tab}.csv').open(newline='', encoding='utf-8') as stream:
            rows = list(csv.DictReader(stream))
        sheet_rows += len(rows)
        by_id = {row['Master Record ID']: row for row in rows}
        assert len(by_id) == len(rows)
        target = [farmer for farmer in farmers if (fulfillment_partner_for_farmer(farmer) == 'ECOCONSERVE') == (tab == 'Eco-conserve')]
        assert set(by_id) == {str(farmer.pk) for farmer in target}
        for farmer in completed:
            if str(farmer.pk) not in by_id:
                continue
            row = by_id[str(farmer.pk)]
            assert row['Case ID'] == display_case_reference(farmer.pk)
            assert row['Current Pipeline State'] == 'COMMISSIONED'
            assert row['Payment No'] == str(batch.payment_number)
            assert Decimal(row['Balance Due']) == farmer.balance_due.quantize(Decimal('1'), rounding=ROUND_HALF_UP)
            action = HomeBiogasAction.objects.get(farmer=farmer)
            assert action.installation_date <= action.commissioning_date <= timezone.localdate()
            assert farmer.customer_no and farmer.credit_decision == farmer.final_decision == 'Approved'
    assert sheet_rows == len(farmers)
    for order in RequisitionBatch.objects.all():
        assert sha256(bytes(order.file_content)).hexdigest() == order.content_checksum
        workbook = load_workbook(BytesIO(bytes(order.file_content)))
        assert workbook.active.cell(9, 2).value == JawabuFarmerMaster.objects.get(pk=order.farmer_ids[0]).customer_name
    manifest = json.loads((output / 'drive' / 'manifest.json').read_text(encoding='utf-8'))
    for entry in manifest:
        content = (output / 'drive' / entry['id']).read_bytes()
        assert sha256(content).hexdigest() == entry['id'] and len(content) == entry['size']
    result = {'status': 'passed', 'intake': len(farmers), 'completed': expected,
              'source': evidence['source'], 'sheet_rows': sheet_rows,
              'orders': RequisitionBatch.objects.count(), 'matched_invoices': expected,
              'commissioned': expected, 'payment_batches': 1,
              'operations': IntegrationOperation.objects.count(), 'drive_uploads': len(manifest),
              'unique_blobs': len({entry['id'] for entry in manifest}),
              'original_harness_status': evidence['status'], 'elapsed_seconds': evidence['elapsed_seconds']}
    (output / 'verification.json').write_text(json.dumps(result, indent=2), encoding='utf-8')
    print(json.dumps(result))
    return result


def resume_sysup(output):
    """Resume an interrupted synthetic review using its saved source and current revisions."""
    from django.conf import settings
    from django.test import Client
    from django.urls import reverse
    from core.models import JawabuFarmerUploadBatch, JawabuFarmerMaster, UserProfile, IntegrationOperation
    from core.tests import FakeMasterDataSheet, FakeJawabuService
    from core.services.portal_publication import queued_publication_operations, attempt_publication

    settings.TELEGRAM_BOT_TOKEN = '000000:synthetic-local-audit-only'
    settings.PORTAL_PUBLICATION_MIN_SPACING_SECONDS = 0
    profile = UserProfile.objects.get(user__username='synthetic-operations_admin')
    fields = {'auth_date': str(int(time.time())), 'user': json.dumps({'id': int(profile.telegram_id)})}
    secret = hmac.new(b'WebAppData', settings.TELEGRAM_BOT_TOKEN.encode(), sha256).digest()
    fields['hash'] = hmac.new(secret, '\n'.join(f'{key}={value}' for key, value in sorted(fields.items())).encode(), sha256).hexdigest()
    client = Client(HTTP_X_TELEGRAM_INIT_DATA=urlencode(fields))
    batch = JawabuFarmerUploadBatch.objects.get(import_kind='system_export')
    original_hash = sha256(bytes(batch.source_content)).hexdigest()
    result = {'initially_committed': batch.committed_count, 'remaining_before': len(batch.parsed_rows),
              'requests': [], 'source_rows': batch.total_rows, 'status': 'running'}
    started = time.monotonic()
    while batch.parsed_rows:
        rows = [{**row, 'approved': True} for row in batch.parsed_rows[:settings.SYSUP_PROCESSING_CHUNK_SIZE]]
        assert all(row.get('Matched Farmer ID') for row in rows)
        payload = {'rows': rows, 'portal_revision': batch.portal_revision,
                   'client_request_id': f'synthetic-resume-{batch.pk}-{batch.portal_revision}'}
        before = time.monotonic()
        response = client.post(reverse('portal_import_commit', args=(batch.pk,)), json.dumps(payload), content_type='application/json')
        data = response.json()
        assert response.status_code == 200 and data['ok'], data
        assert 0 < data['result']['committed'] <= settings.SYSUP_COMMIT_MAX_ROWS_PER_REQUEST, data
        result['requests'].append({'selected': len(rows), 'committed': data['result']['committed'],
                                   'duration_ms': round((time.monotonic() - before) * 1000, 2)})
        batch.refresh_from_db()
        (output / 'sysup-resume.json').write_text(json.dumps(result, indent=2), encoding='utf-8')
        print('resume', batch.committed_count, 'of', batch.total_rows, flush=True)
    assert batch.status == 'committed' and batch.committed_count == batch.total_rows
    assert sha256(bytes(batch.source_content)).hexdigest() == original_hash
    assert JawabuFarmerMaster.objects.exclude(customer_no='').count() == batch.total_rows
    sheet_dir = output / 'sheets'
    sheet_dir.mkdir(exist_ok=True)
    with (output.parent / 'sheets' / 'Master Data.csv').open(newline='', encoding='utf-8') as stream:
        headers = next(csv.reader(stream))
    sheets = {tab: FakeMasterDataSheet(headers) for tab in ('Master Data', 'Eco-conserve', 'Orders')}
    for sheet in sheets.values():
        sheet.values = [list(headers)]
    def service(*, sheet_id, sheet_name, **kwargs):
        assert sheet_id == 'local-audit-sheet'
        return FakeJawabuService(sheets[sheet_name])
    with patch('core.services.sheets.GoogleSheetsService.get_instance', side_effect=service):
        attempts = 0
        while (operation := queued_publication_operations().first()) is not None:
            attempted = attempt_publication(operation)
            assert not attempted.get('error'), attempted
            attempts += 1
            assert attempts <= batch.total_rows * 3
    assert not IntegrationOperation.objects.exclude(status='succeeded').exists()
    for tab, sheet in sheets.items():
        with (sheet_dir / f'{tab}.csv').open('w', newline='', encoding='utf-8') as stream:
            csv.writer(stream).writerows(sheet.values)
    assert sum(len(sheet.values) - 1 for sheet in sheets.values()) == batch.total_rows
    result.update(status='passed', final_committed=batch.committed_count, remaining_after=0,
                  sheet_rows=batch.total_rows, publication_attempts=attempts,
                  elapsed_seconds=round(time.monotonic() - started, 2))
    (output / 'sysup-resume.json').write_text(json.dumps(result, indent=2), encoding='utf-8')
    print(json.dumps({key: value for key, value in result.items() if key != 'requests'}))
    return result
