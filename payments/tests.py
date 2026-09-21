from decimal import Decimal
import json
from unittest.mock import patch

from django.apps import apps
from django.contrib.auth import get_user_model
from django.test import RequestFactory, TestCase
from django.urls import reverse
from openpyxl import Workbook

from core.models import GroupSheetConfiguration, InvoiceUploadBatch, JawabuFarmerMaster, ParsedInvoice, PaymentDocument
from core.services.payment_documents import _write_payment_mode
from core.services.jawabu_case_reference import display_case_reference
from core.services.jawabu_pipeline import completed_payment_number_for_farmer
from payments.models import PaymentBatch, PaymentCaseReview, PaymentReceiptItem, PaymentSequenceState
from payments.receipt_batches import create_payment_batch_from_receipt, create_receipt_batch
from payments.services import (
    PaymentBatchError,
    add_cases,
    adjust_sequence,
    cancel_batch,
    complete_batch_for_document,
    create_batch,
    generate_reviewed_workbook,
    remove_case,
    review_case,
    serialize_batch,
    submit_for_review,
    update_case_mode,
)


class PaymentBatchServiceTests(TestCase):
    def setUp(self):
        self.user = get_user_model().objects.create_user(username='payment-operator')
        self.group = GroupSheetConfiguration.objects.create(
            group_id='-100-payment-test', display_name='Payment test', enabled=True,
            sheet_id='test-sheet', workflow={'type': 'jawabu_homebiogas'},
        )

    def farmer(self, suffix):
        return JawabuFarmerMaster.objects.create(
            customer_name=f'Customer {suffix}', national_id=f'12345{suffix}',
            primary_phone=f'25470000{suffix}', customer_no=f'CUST-{suffix}',
            order_number=f'ORDER-{suffix}', invoice_number=f'INV-{suffix}',
            balance_due=Decimal('1000'), repayment_date='10TH', repayment_day=10,
            repayment_tenor='6', repayment_tenor_months=6,
            payment_product='Business', final_decision='Approved', workflow_revision=1,
        )

    def batch(self, **kwargs):
        return create_batch(group_configuration=self.group, **kwargs)

    def add(self, batch, *farmers, mode='CASH', modes=None):
        payment_modes = modes or {str(farmer.id): mode for farmer in farmers}
        return add_cases(
            batch.id, farmer_ids=[farmer.id for farmer in farmers], payment_modes=payment_modes,
            expected_revision=batch.revision,
        )

    @staticmethod
    def ready(*args, farmer_ids=None, **kwargs):
        return {
            'ready': [{'farmer_id': value, 'row': {}} for value in (farmer_ids or [])],
            'blocked': [], 'ready_count': len(farmer_ids or []), 'blocked_count': 0,
        }

    @patch('payments.services.payment_readiness', side_effect=ready.__func__)
    def test_modes_are_governed(self, _readiness):
        farmer = self.farmer('00')
        batch = self.batch()
        batch = self.add(batch, farmer, mode='CASH')
        self.assertEqual(batch.case_memberships.get(farmer=farmer).payment_mode, 'CASH')
        with self.assertRaisesMessage(PaymentBatchError, 'Choose Loan - Jawabu or Cash'):
            add_cases(
                batch.id, farmer_ids=[farmer.id], payment_modes={str(farmer.id): 'CHEQUE'},
                expected_revision=batch.revision,
            )

    @patch('payments.services.payment_readiness', side_effect=ready.__func__)
    def test_one_batch_keeps_an_independent_mode_for_each_case(self, _readiness):
        cash, loan = self.farmer('14'), self.farmer('15')
        batch = self.batch()
        batch = self.add(
            batch, cash, loan,
            modes={str(cash.id): 'CASH', str(loan.id): 'LOAN-JAWABU'},
        )

        payload = serialize_batch(batch)
        modes = {item['farmer_id']: item['payment_mode'] for item in payload['cases']}
        self.assertEqual(modes, {str(cash.id): 'CASH', str(loan.id): 'LOAN-JAWABU'})
        self.assertEqual(
            {item['farmer_id']: item['case_reference'] for item in payload['cases']},
            {str(cash.id): display_case_reference(cash), str(loan.id): display_case_reference(loan)},
        )
        self.assertEqual(payload['payment_mode_counts'], {'LOAN-JAWABU': 1, 'CASH': 1})

    @patch('payments.services.payment_readiness', side_effect=ready.__func__)
    def test_batch_payload_exposes_the_case_facts_needed_for_payment_review(self, _readiness):
        farmer = self.farmer('identity')
        farmer.branch = 'Embu'
        farmer.system_branch = 'Embu Central'
        farmer.jbl_officer = 'Legacy BRO'
        farmer.system_loan_officer = 'Current BRO'
        farmer.save(update_fields=['branch', 'system_branch', 'jbl_officer', 'system_loan_officer', 'updated_at'])

        payload = serialize_batch(self.add(self.batch(), farmer))
        case = payload['cases'][0]

        self.assertEqual(case['primary_phone'], farmer.primary_phone)
        self.assertEqual(case['branch'], 'Embu Central')
        self.assertEqual(case['loan_officer'], 'Current BRO')

    @patch('payments.services.payment_readiness', side_effect=ready.__func__)
    def test_batch_payload_exposes_ordinal_repayment_day(self, _readiness):
        farmer = self.farmer('18')
        farmer.repayment_date = '2026-09-02'
        farmer.repayment_day = None
        farmer.save(update_fields=['repayment_date', 'repayment_day', 'updated_at'])
        batch = self.add(self.batch(), farmer)

        self.assertEqual(serialize_batch(batch)['cases'][0]['preferred_repayment_date'], '2ND')

    @patch('payments.services.payment_readiness', side_effect=ready.__func__)
    def test_generation_receives_the_exact_mode_for_every_approved_case(self, _readiness):
        cash, loan = self.farmer('16'), self.farmer('17')
        batch = self.add(
            self.batch(), cash, loan,
            modes={str(cash.id): 'CASH', str(loan.id): 'LOAN-JAWABU'},
        )
        batch = submit_for_review(batch.id, expected_revision=batch.revision)
        for farmer in (cash, loan):
            batch = review_case(
                batch.id, farmer.id, decision='approved', comment='Confirmed.',
                expected_revision=batch.revision, actor=self.user,
            )
        document = PaymentDocument.objects.create(
            order_number='PAYMENT-1', payment_number='1', status='awaiting_scan', version=1,
        )
        with patch('payments.services.create_payment_document', return_value=document) as generator:
            generate_reviewed_workbook(batch.id, expected_revision=batch.revision, actor=self.user)
        self.assertEqual(generator.call_args.kwargs['case_payment_modes'], {
            str(cash.id): 'CASH', str(loan.id): 'LOAN-JAWABU',
        })

    def test_request_key_replays_only_the_same_payment_meaning(self):
        batch = self.batch(request_id='create-payment-1')
        replay = self.batch(request_id='create-payment-1')
        self.assertEqual(replay.pk, batch.pk)
        with self.assertRaisesMessage(PaymentBatchError, 'different payment change'):
            create_batch(group_configuration=GroupSheetConfiguration.objects.create(
                group_id='-100-other-payment-test', display_name='Other payment', enabled=True,
                sheet_id='other-sheet', workflow={'type': 'jawabu_homebiogas'},
            ), request_id='create-payment-1')

    @patch('payments.services.payment_readiness', side_effect=ready.__func__)
    def test_payment_numbers_allocate_consecutively_and_retry_idempotently(self, _readiness):
        first_farmer, second_farmer = self.farmer('01'), self.farmer('02')
        first = self.batch()
        first = self.add(first, first_farmer, mode='LOAN-JAWABU')
        prior_revision = first.revision
        first = submit_for_review(first.id, expected_revision=prior_revision, actor=self.user, request_id='submit-1')
        replay = submit_for_review(first.id, expected_revision=prior_revision, actor=self.user, request_id='submit-1')
        self.assertIsNone(first.payment_number)
        self.assertIsNone(replay.payment_number)
        first = review_case(
            first.id, first_farmer.id, decision='approved', comment='Ready.',
            expected_revision=first.revision, actor=self.user,
        )
        first_generation_revision = first.revision

        def generated_document(order_number, payment_number, **_kwargs):
            return PaymentDocument.objects.create(
                order_number=order_number, payment_number=payment_number,
                status='awaiting_scan', version=1,
            )

    @patch('payments.services.payment_readiness', side_effect=ready.__func__)
    @patch('payments.receipt_batches._item_disposition')
    def test_invoice_delivery_creates_only_its_reconciled_payment_rows(self, disposition, _readiness):
        payable = self.farmer('receipt-payable')
        held = self.farmer('receipt-held')
        upload = InvoiceUploadBatch.objects.create(
            original_filename='HB invoices.pdf', status='matched', total_pages=2, total_parsed=2,
        )
        parsed_payable = ParsedInvoice.objects.create(
            batch=upload, page=1, invoice_no='INV-R1', customer_name=payable.customer_name,
            customer_id=payable.national_id, status='matched', matched_farmer=payable,
        )
        parsed_held = ParsedInvoice.objects.create(
            batch=upload, page=2, invoice_no='INV-R2', customer_name=held.customer_name,
            customer_id='different-id', status='matched', matched_farmer=held,
        )

        disposition.side_effect = [
            (PaymentReceiptItem.STATUS_MATCHED, payable, ''),
            (PaymentReceiptItem.STATUS_NAME_CHANGE, held, 'Invoice holder differs from the committed loan applicant.'),
        ]
        receipt, replayed = create_receipt_batch(
            group_configuration=self.group, uploads=[upload], actor=self.user, request_id='receipt-delivery-1',
        )
        self.assertFalse(replayed)
        self.assertEqual(
            set(receipt.items.values_list('invoice_id', 'status')),
            {(parsed_payable.id, PaymentReceiptItem.STATUS_MATCHED), (parsed_held.id, PaymentReceiptItem.STATUS_NAME_CHANGE)},
        )

        batch, replayed = create_payment_batch_from_receipt(
            receipt_id=receipt.id, expected_revision=receipt.revision,
            payment_modes={str(payable.id): 'LOAN-JAWABU'}, actor=self.user, request_id='receipt-payment-1',
        )
        self.assertFalse(replayed)
        self.assertEqual(list(batch.case_memberships.filter(is_active=True).values_list('farmer_id', flat=True)), [payable.id])
        payload = serialize_batch(batch)
        self.assertEqual(payload['receipt_batch_id'], str(receipt.id))
        self.assertEqual(payload['held_items'][0]['farmer_id'], str(held.id))

        unrelated = self.farmer('receipt-unrelated')
        with self.assertRaisesMessage(PaymentBatchError, 'only cases reconciled in that delivery'):
            add_cases(
                batch.id, farmer_ids=[unrelated.id], payment_modes={str(unrelated.id): 'CASH'},
                expected_revision=batch.revision,
            )

        with patch('payments.services.create_payment_document', side_effect=generated_document):
            first = generate_reviewed_workbook(
                first.id, expected_revision=first_generation_revision,
                actor=self.user, request_id='generate-1',
            )
            replay = generate_reviewed_workbook(
                first.id, expected_revision=first_generation_revision,
                actor=self.user, request_id='generate-1',
            )
        second = self.batch()
        second = self.add(second, second_farmer)
        second = submit_for_review(second.id, expected_revision=second.revision, actor=self.user, request_id='submit-2')
        second = review_case(
            second.id, second_farmer.id, decision='approved', comment='Ready.',
            expected_revision=second.revision, actor=self.user,
        )
        with patch('payments.services.create_payment_document', side_effect=generated_document):
            second = generate_reviewed_workbook(
                second.id, expected_revision=second.revision,
                actor=self.user, request_id='generate-2',
            )
        self.assertEqual((first.payment_number, replay.payment_number, second.payment_number), (1, 1, 2))
        self.assertEqual(PaymentSequenceState.objects.get(group_configuration=self.group).next_number, 3)

    @patch('payments.services.payment_readiness', side_effect=ready.__func__)
    def test_review_progress_is_durable_and_changed_values_invalidate_only_that_case(self, _readiness):
        first_farmer, second_farmer = self.farmer('03'), self.farmer('04')
        batch = self.batch()
        batch = self.add(batch, first_farmer, second_farmer, mode='LOAN-JAWABU')
        batch = submit_for_review(batch.id, expected_revision=batch.revision)
        batch = review_case(
            batch.id, first_farmer.id, decision='approved', comment='Payment details confirmed.',
            expected_revision=batch.revision, actor=self.user,
        )
        payload = serialize_batch(batch)
        self.assertEqual(payload['counts']['approved'], 1)
        self.assertEqual(payload['counts']['pending'], 1)
        first_farmer.balance_due = Decimal('1200')
        first_farmer.save(update_fields=['balance_due', 'updated_at'])
        payload = serialize_batch(PaymentBatch.objects.get(pk=batch.pk))
        self.assertEqual(payload['counts']['approved'], 0)
        self.assertEqual(payload['counts']['changed'], 1)

    @patch('payments.services.payment_readiness', side_effect=ready.__func__)
    def test_membership_change_before_scan_supersedes_workbook(self, _readiness):
        farmer = self.farmer('05')
        batch = self.batch()
        batch = self.add(batch, farmer)
        document = PaymentDocument.objects.create(
            order_number='PAYMENT-1', payment_number='1', status='awaiting_scan', version=1,
        )
        batch.payment_number = 1
        batch.status = PaymentBatch.STATUS_AWAITING_SCAN
        batch.current_document = document
        batch.save()
        batch = remove_case(
            batch.id, farmer.id, reason='Customer moved to a later payment.',
            expected_revision=batch.revision, actor=self.user,
        )
        document.refresh_from_db()
        self.assertEqual(document.status, 'superseded')
        self.assertIsNone(batch.current_document_id)
        self.assertFalse(batch.case_memberships.get(farmer=farmer).is_active)

    @patch('payments.services.payment_readiness', side_effect=ready.__func__)
    def test_signed_scan_completion_locks_membership(self, _readiness):
        farmer = self.farmer('06')
        batch = self.batch()
        batch = self.add(batch, farmer, mode='LOAN-JAWABU')
        document = PaymentDocument.objects.create(
            order_number='PAYMENT-1', payment_number='1', status='awaiting_scan', version=1,
        )
        batch.payment_number = 1
        batch.status = PaymentBatch.STATUS_AWAITING_SCAN
        batch.current_document = document
        batch.save()
        batch = complete_batch_for_document(document, actor=self.user)
        self.assertEqual(batch.status, PaymentBatch.STATUS_COMPLETED)
        document.refresh_from_db()
        self.assertEqual(document.status, 'completed')
        self.assertEqual(completed_payment_number_for_farmer(farmer), '1')
        farmer.refresh_from_db()
        self.assertEqual(farmer.workflow_revision, 2)
        with self.assertRaisesMessage(PaymentBatchError, 'cannot be removed'):
            remove_case(batch.id, farmer.id, reason='Too late', expected_revision=batch.revision)

    @patch('payments.services.payment_readiness', side_effect=ready.__func__)
    def test_remove_requires_reason(self, _readiness):
        farmer = self.farmer('07')
        batch = self.batch()
        batch = self.add(batch, farmer)
        with self.assertRaisesMessage(PaymentBatchError, 'Give a reason'):
            remove_case(batch.id, farmer.id, reason='', expected_revision=batch.revision)

    @patch('payments.services.payment_readiness', side_effect=ready.__func__)
    def test_cancelling_releases_case_but_keeps_membership_history(self, _readiness):
        farmer = self.farmer('08')
        first = self.batch()
        first = self.add(first, farmer)
        first = cancel_batch(
            first.id, reason='Rebuild with a different set of clients.',
            expected_revision=first.revision, actor=self.user,
        )
        second = self.batch()
        second = self.add(second, farmer, mode='LOAN-JAWABU')
        self.assertEqual(first.status, PaymentBatch.STATUS_CANCELLED)
        self.assertTrue(first.case_memberships.get(farmer=farmer).is_active)
        self.assertTrue(second.case_memberships.get(farmer=farmer).is_active)

    @patch('payments.services.payment_readiness', side_effect=ready.__func__)
    def test_sequence_adjustment_cannot_reuse_an_allocated_number(self, _readiness):
        farmer = self.farmer('09')
        batch = self.batch()
        batch = self.add(batch, farmer)
        batch = submit_for_review(batch.id, expected_revision=batch.revision)
        batch = review_case(
            batch.id, farmer.id, decision='approved', comment='Ready.',
            expected_revision=batch.revision, actor=self.user,
        )
        document = PaymentDocument.objects.create(
            order_number='PAYMENT-1', payment_number='1', status='awaiting_scan', version=1,
        )
        with patch('payments.services.create_payment_document', return_value=document):
            batch = generate_reviewed_workbook(batch.id, expected_revision=batch.revision, actor=self.user)
        self.assertEqual(batch.payment_number, 1)
        with self.assertRaisesMessage(PaymentBatchError, 'higher than the allocated number 1'):
            adjust_sequence(group_configuration=self.group, next_number=1, reason='Unsafe rollback', actor=self.user)
        state = adjust_sequence(
            group_configuration=self.group, next_number=20,
            reason='Continue from the finance register.', actor=self.user, request_id='adjust-20',
        )
        replay = adjust_sequence(
            group_configuration=self.group, next_number=20,
            reason='Continue from the finance register.', actor=self.user, request_id='adjust-20',
        )
        self.assertEqual((state.next_number, replay.next_number), (20, 20))
        with self.assertRaisesMessage(PaymentBatchError, 'already used for a different'):
            adjust_sequence(
                group_configuration=self.group, next_number=21,
                reason='A different operation.', actor=self.user, request_id='adjust-20',
            )
        with self.assertRaisesMessage(PaymentBatchError, 'changed while you were working'):
            adjust_sequence(
                group_configuration=self.group, next_number=21,
                reason='Valid new number, stale screen.', actor=self.user, expected_revision=1,
            )

    @patch('payments.services.payment_readiness', side_effect=ready.__func__)
    def test_mode_or_membership_change_invalidates_existing_case_approvals(self, _readiness):
        first_farmer, second_farmer = self.farmer('11'), self.farmer('12')
        batch = self.batch()
        batch = self.add(batch, first_farmer)
        batch = submit_for_review(batch.id, expected_revision=batch.revision)
        batch = review_case(
            batch.id, first_farmer.id, decision='approved', comment='Confirmed.',
            expected_revision=batch.revision, actor=self.user,
        )
        batch = update_case_mode(
            batch.id, first_farmer.id, payment_mode='LOAN-JAWABU',
            expected_revision=batch.revision, actor=self.user,
        )
        self.assertEqual(batch.case_memberships.get(farmer=first_farmer).review.decision, 'pending')

        batch = review_case(
            batch.id, first_farmer.id, decision='approved', comment='Confirmed again.',
            expected_revision=batch.revision, actor=self.user,
        )
        batch = self.add(batch, second_farmer)
        self.assertEqual(batch.case_memberships.get(farmer=first_farmer).review.decision, 'pending')

    @patch('payments.services.payment_readiness', side_effect=ready.__func__)
    def test_changed_payment_values_block_generation_and_reopen_review(self, _readiness):
        farmer = self.farmer('13')
        batch = self.batch()
        batch = self.add(batch, farmer)
        batch = submit_for_review(batch.id, expected_revision=batch.revision)
        batch = review_case(
            batch.id, farmer.id, decision='approved', comment='Ready.',
            expected_revision=batch.revision, actor=self.user,
        )
        farmer.repayment_date = '15TH'
        farmer.save(update_fields=['repayment_date', 'updated_at'])
        with patch('payments.services.create_payment_document') as generator:
            with self.assertRaisesMessage(PaymentBatchError, 'changed after review'):
                generate_reviewed_workbook(batch.id, expected_revision=batch.revision, actor=self.user)
        generator.assert_not_called()
        batch.refresh_from_db()
        self.assertEqual(batch.status, PaymentBatch.STATUS_IN_REVIEW)
        self.assertEqual(batch.case_memberships.get(farmer=farmer).review.decision, 'pending')

    @patch('payments.services.payment_readiness', side_effect=ready.__func__)
    def test_batch_change_during_generation_supersedes_generated_workbook(self, _readiness):
        farmer = self.farmer('10')
        batch = self.batch()
        batch = self.add(batch, farmer)
        batch = submit_for_review(batch.id, expected_revision=batch.revision)
        batch = review_case(
            batch.id, farmer.id, decision='approved', comment='Ready for payment.',
            expected_revision=batch.revision, actor=self.user,
        )
        document = PaymentDocument.objects.create(
            order_number='PAYMENT-1', payment_number='1', status='awaiting_scan', version=1,
        )

        def mutate_during_generation(*args, **kwargs):
            current = PaymentBatch.objects.get(pk=batch.pk)
            PaymentBatch.objects.filter(pk=batch.pk).update(revision=current.revision + 1)
            return document

        with patch('payments.services.create_payment_document', side_effect=mutate_during_generation):
            with self.assertRaisesMessage(PaymentBatchError, 'changed while the workbook was being generated'):
                generate_reviewed_workbook(batch.id, expected_revision=batch.revision, actor=self.user)
        document.refresh_from_db()
        self.assertEqual(document.status, 'superseded')


class PaymentWorkflowContractTests(TestCase):
    def test_payment_error_uses_current_message_contract_with_actionable_copy(self):
        from core.api.portal_views import _portal_payment_batch_error

        request = RequestFactory().post(
            '/api/portal/payments/batches/example/generate/',
            HTTP_X_MINIAPP_MESSAGE_CONTRACT='2',
        )
        response = _portal_payment_batch_error(
            request,
            PaymentBatchError(
                'The payment workbook template needs attention.',
                code='payment_workbook_template_invalid',
            ),
        )
        payload = json.loads(response.content)

        self.assertEqual(response.status_code, 400)
        self.assertEqual(payload['code'], 'payment_workbook_template_invalid')
        self.assertEqual(payload['message'], 'The payment workbook template needs attention.')
        self.assertNotIn('error', payload)

    def test_payment_mode_marker_supports_configured_and_detected_cells(self):
        configured = Workbook()
        ws = configured.active
        self.assertTrue(_write_payment_mode(ws, {'payment_mode_cell': 'D4'}, 'CASH'))
        self.assertEqual(ws['D4'].value, 'CASH')

        detected = Workbook()
        ws = detected.active
        ws['B2'] = 'Mode of Payment'
        self.assertTrue(_write_payment_mode(ws, {}, 'LOAN-JAWABU'))
        self.assertEqual(ws['C2'].value, 'LOAN-JAWABU')

    def test_payment_models_and_fields_are_documented(self):
        for model in apps.get_app_config('payments').get_models():
            self.assertTrue(model._meta.db_table_comment, model._meta.label)
            for field in model._meta.concrete_fields:
                self.assertTrue(field.db_comment, f'{model._meta.label}.{field.name}')

    def test_governed_payment_routes_are_named(self):
        batch_id = '00000000-0000-0000-0000-000000000001'
        farmer_id = '00000000-0000-0000-0000-000000000002'
        self.assertEqual(reverse('portal_payment_batches'), '/api/portal/payments/batches/')
        self.assertEqual(reverse('portal_payment_sequence'), '/api/portal/payments/sequence/')
        self.assertIn(batch_id, reverse('portal_payment_batch_detail', kwargs={'batch_id': batch_id}))
        self.assertIn(farmer_id, reverse(
            'portal_payment_batch_case_review', kwargs={'batch_id': batch_id, 'farmer_id': farmer_id},
        ))
        self.assertIn(farmer_id, reverse(
            'portal_payment_batch_case_mode', kwargs={'batch_id': batch_id, 'farmer_id': farmer_id},
        ))
