"""Synthetic one-off disclosure and email presentation regressions."""
import base64
import uuid
from io import BytesIO
from unittest.mock import patch

from django.contrib.auth import get_user_model
from django.core.exceptions import PermissionDenied, ValidationError
from django.test import SimpleTestCase, TestCase, override_settings
from openpyxl import load_workbook

from core.models import AccessGrant, GroupSheetConfiguration, JawabuFarmerMaster
from .models import ApprovedRecipient, ReportDelivery, ReportSchedule, WebhookReceipt
from .one_off import queue_export, validate_export, capture_export
from .services import claim_delivery, process_delivery, reconcile_webhooks
from .rendering import build_payload
from .views import report_export_response


class EmailPresentationTests(SimpleTestCase):
    def snapshot(self, **changes):
        return {'preset': 'complaints', 'summary': {'Complaints received': 12, 'Open': 6, 'Resolved': 6, 'Needs details': 0},
                'charts': [], 'rows': [], 'exported_rows': 12, 'total_rows': 12,
                'xlsx_content': base64.b64encode(b'synthetic-workbook').decode(),
                'applied_filters': {}, 'period': {'from': '2026-10-01', 'to': '2026-10-05'},
                'run_at': '2026-10-05T09:00:00+03:00', **changes}

    def test_logo_is_local_inline_and_excel_bytes_are_unchanged(self):
        from .rendering import BRAND_LOGO_PATH
        snapshot = self.snapshot()
        payload = build_payload(snapshot, 'management@example.invalid', {})
        self.assertEqual(payload['attachments'][0]['content'], snapshot['xlsx_content'])
        logo = payload['attachments'][1]
        self.assertEqual(base64.b64decode(logo['content']), BRAND_LOGO_PATH.read_bytes())
        self.assertEqual(logo['content_type'], 'image/png')
        self.assertIn('src="cid:' + logo['content_id'] + '"', payload['html'])
        self.assertNotIn('https://', payload['html'])
        self.assertNotIn('report_download_url', payload['html'])

    def test_missing_logo_keeps_delivery_and_text_brand(self):
        with patch('report_delivery.rendering.BRAND_LOGO_PATH') as logo:
            logo.read_bytes.side_effect = FileNotFoundError
            payload = build_payload(self.snapshot(), 'management@example.invalid', {})
        self.assertEqual(len(payload['attachments']), 1)
        self.assertNotIn('src="cid:', payload['html'])
        self.assertIn('JAWABU BIASHARA', payload['html'])

    def test_readable_duration_percent_zero_and_missing_values_without_mutation(self):
        from copy import deepcopy
        from .rendering import email_context
        snapshot = self.snapshot(summary={'Median resolution (hours)': 62, 'Median HB response (hours)': 0.25,
                                          'Median TAT (minutes)': 0, 'Resolved on time (%)': 87.5,
                                          'Missing timing (hours)': None, 'Open': 0})
        before = deepcopy(snapshot)
        context = email_context(snapshot)
        self.assertEqual([m['value'] for m in context['metrics']], ['2 days 14 hr', '15 min', '0 min', '87.50%', '—', '0'])
        self.assertEqual(context['metrics'][0]['label'], 'Median resolution')
        self.assertEqual(snapshot, before)
        self.assertEqual(email_context(self.snapshot(summary={'Median resolution (hours)': 24.25}))['metrics'][0]['value'], '1 day 15 min')

    def test_breakdowns_pair_small_tables_but_keep_multi_series_full_width(self):
        from .rendering import email_context
        charts = [{'title': title, 'labels': ['Training branch'], 'datasets': [{'label': 'Cases', 'values': [3]}]}
                  for title in ['Categories', 'Branches']]
        charts.append({'title': 'Finance', 'labels': ['Training branch'],
                       'datasets': [{'label': str(i), 'values': [i]} for i in range(4)]})
        context = email_context(self.snapshot(charts=charts))
        self.assertEqual([len(group) for group in context['breakdown_groups']], [2, 1])
        self.assertEqual(len(context['breakdown_groups'][1][0]['headers']), 4)

    def test_empty_report_and_long_scope_are_escaped(self):
        payload = build_payload(self.snapshot(summary={}, total_rows=0, exported_rows=0,
                                             applied_filters={'branch': '<script>Training</script>'}),
                                'management@example.invalid', {})
        self.assertIn('0 of 0 matching cases', payload['html'])
        self.assertIn('&lt;script&gt;Training&lt;/script&gt;', payload['html'])
        self.assertNotIn('<script>', payload['html'])

    def chart(self, **changes):
        return {'title': 'Complaints by branch', 'labels': ['Training branch', 'Second branch'],
                'datasets': [{'label': 'Complaints', 'values': [5, 0]}], **changes}

    def test_larger_typography_and_full_width_amounts(self):
        from .rendering import email_context
        snapshot = self.snapshot(summary={'Cases': 12, 'Invoice total (KES)': '123456789'})
        context = email_context(snapshot)
        self.assertTrue(context['roomy_headline'])
        self.assertEqual(context['headline_width'], 50)
        payload = build_payload(snapshot, 'management@example.invalid', {})
        self.assertIn("'Segoe UI','Helvetica Neue',Arial,sans-serif", payload['html'])
        self.assertNotIn('font-size:12px', payload['html'])
        self.assertIn('font-size:30px', payload['html'])
        self.assertIn('123,456,789', payload['html'])

    def test_charts_are_offline_inline_pngs_and_snapshot_is_unchanged(self):
        from copy import deepcopy
        from io import BytesIO
        from PIL import Image
        snapshot = self.snapshot(charts=[self.chart()])
        before = deepcopy(snapshot)
        payload = build_payload(snapshot, 'management@example.invalid', {})
        graph = payload['attachments'][2]
        self.assertEqual(graph['content_type'], 'image/png')
        self.assertIn('src="cid:' + graph['content_id'] + '"', payload['html'])
        self.assertIn('Values are listed in the breakdown below.', payload['html'])
        with Image.open(BytesIO(base64.b64decode(graph['content']))) as image:
            self.assertEqual(image.format, 'PNG')
            self.assertEqual(image.width, 600)
        self.assertEqual(payload['attachments'][0]['content'], before['xlsx_content'])
        self.assertEqual(snapshot, before)
        self.assertNotIn('https://', payload['html'])
        self.assertNotIn('<script', payload['html'])

    def test_only_first_three_nonempty_charts_are_considered(self):
        charts = [self.chart(labels=[])] + [self.chart(title=f'Graph {i}') for i in range(5)]
        payload = build_payload(self.snapshot(charts=charts), 'management@example.invalid', {})
        graphs = [a for a in payload['attachments'] if a.get('content_id', '').startswith('jawabu-report-chart')]
        self.assertEqual(len(graphs), 3)
        self.assertEqual(len({a['content_id'] for a in graphs}), 3)
        self.assertIn('Graph 4', payload['html'])  # Tables remain for the remaining graphs.

    def test_dense_chart_retains_every_group_in_html_fallback(self):
        labels = [f'Training group {i}' for i in range(12)]
        payload = build_payload(self.snapshot(charts=[self.chart(labels=labels,
                                     datasets=[{'label': 'Cases', 'values': list(range(12))}])]),
                                'management@example.invalid', {})
        self.assertEqual(len(payload['attachments']), 2)
        for label in labels:
            self.assertIn(label, payload['html'])
        self.assertNotIn('First 8 groups', payload['html'])

    def test_chart_render_failure_keeps_report_tables_and_workbook(self):
        with patch('report_delivery.email_charts.chart_png', side_effect=OSError('synthetic font failure')):
            payload = build_payload(self.snapshot(charts=[self.chart()]), 'management@example.invalid', {})
        self.assertEqual(len(payload['attachments']), 2)
        self.assertIn('Training branch', payload['html'])
        self.assertNotIn('class="report-graph"', payload['html'])

    def test_png_zero_negative_sparse_line_and_multi_series(self):
        from .email_charts import chart_png
        for chart in [self.chart(), self.chart(datasets=[{'label': 'Cases', 'values': [-2, 4]}]),
                      self.chart(type='line', labels=['01 Oct', '02 Oct'],
                                 datasets=[{'label': 'Cases', 'values': [0, 2]}]),
                      self.chart(datasets=[{'label': 'Cases', 'values': [5]}]),
                      self.chart(datasets=[{'label': 'Cases', 'values': [0, 0]}]),
                      self.chart(datasets=[{'label': 'Cases', 'values': [5, None]}]),
                      self.chart(datasets=[{'label': str(i), 'values': [5, 0]} for i in range(4)])]:
            with self.subTest(chart=chart):
                self.assertTrue(chart_png(chart).startswith(b'\x89PNG'))

    def test_unreadable_or_nonfinite_charts_use_tables(self):
        from .email_charts import chart_png
        for chart in [self.chart(labels=['x' * 200]), self.chart(type='pie'),
                      self.chart(datasets=[{'label': 'Cases', 'values': [None, None]}]),
                      self.chart(datasets=[{'label': 'Cases', 'values': ['NaN', 0]}]),
                      self.chart(datasets=[{'label': str(i), 'values': [5, 0]} for i in range(5)])]:
            with self.subTest(chart=chart):
                self.assertIsNone(chart_png(chart))

    def test_chart_bytes_are_included_in_existing_payload_limit(self):
        from .rendering import ReportTooLarge
        with patch('report_delivery.rendering.MAX_PAYLOAD_BYTES', 1024):
            with self.assertRaises(ReportTooLarge):
                build_payload(self.snapshot(charts=[self.chart()]), 'management@example.invalid', {})


class ExportSummaryTests(SimpleTestCase):
    def test_tat_summary_matches_current_or_completed_action_view(self):
        from types import SimpleNamespace
        for view, metrics, basis, label in [
            ('current', {'active': 3, 'stalled': 1}, '', 'Active cases'),
            ('performance', {'created': 3, 'finished': 4}, 'completed_stage_actions', 'Completed actions'),
        ]:
            delivery = SimpleNamespace(requested_by=object(), pk=uuid.uuid4(), configuration={
                'workflow':'tat_tracker', 'preset':'tat', 'filters':{'view':view}})
            with patch('report_delivery.one_off.user_access', return_value={}), patch('report_delivery.one_off.has_capability', return_value=False), \
                 patch('core.services.tat_reporting.export_report_xlsx', return_value=(b'xlsx', 4)), \
                 patch('core.services.tat_reporting.report_summary', return_value={'metrics':metrics, 'metric_basis':basis}):
                snapshot = capture_export(delivery)
            self.assertEqual(snapshot['summary'][label], 3 if view == 'current' else 4)

    def test_open_ended_current_period_is_not_labelled_completed(self):
        from .rendering import email_context
        context = email_context({'preset':'tat', 'summary':{}, 'total_rows':0, 'rows':[],
                                 'applied_filters':{'date_from':'2026-10-01','view':'current'},
                                 'run_at':'2026-10-05T09:00:00+03:00'})
        self.assertEqual(context['period'], 'From 01-Oct-2026')
        self.assertIn('Current workload', context['basis'])


@override_settings(REPORT_EMAIL_DELIVERY_ENABLED=True, RESEND_API_KEY='synthetic-key', REPORT_EMAIL_FROM='reports@example.invalid')
class OneOffTests(TestCase):
    def setUp(self):
        self.actor = get_user_model().objects.create_user(username='synthetic-exporter', is_superuser=True)
        self.denied = get_user_model().objects.create_user(username='synthetic-denied')
        self.group = GroupSheetConfiguration.objects.create(group_id='synthetic-export-group')
        self.payload = {'email': 'management@example.invalid', 'client_request_id': str(uuid.uuid4()),
                        'preset': 'pipeline', 'filters': {'search': 'Synthetic matching', 'chart_key': 'branch',
                                                          'bucket_key': 'Training', 'series_key': 'Cases'}}

    def queue(self, workflow='jawabu_portal', payload=None):
        return queue_export(self.actor, workflow, payload or self.payload)

    def test_one_off_is_idempotent_without_schedule_or_approval(self):
        first = self.queue()
        self.assertEqual(first.pk, self.queue().pk)
        self.assertEqual(ReportDelivery.objects.count(), 1)
        self.assertFalse(ReportSchedule.objects.exists())
        self.assertFalse(ApprovedRecipient.objects.exists())
        self.assertIsNone(first.schedule_id)
        self.assertEqual(first.configuration['filters']['bucket_key'], 'Training')

    def test_invalid_email_wrong_preset_and_unauthorized_export_are_rejected(self):
        with self.assertRaises(PermissionDenied):
            queue_export(self.denied, 'jawabu_portal', self.payload)
        for changed in ({'email': 'invalid'}, {'preset': 'tat'}):
            with self.assertRaises(ValidationError):
                self.queue(payload={**self.payload, **changed})

    def test_retry_key_cannot_change_destination_or_filters(self):
        self.queue()
        for changed in ({'email': 'other@example.invalid'}, {'filters': {}}):
            with self.assertRaises(ValidationError):
                self.queue(payload={**self.payload, **changed})

    def test_revoked_authority_blocks_frozen_payload(self):
        delivery = self.queue()
        self.actor.is_superuser = False
        self.actor.save()
        with patch('report_delivery.resend.send_email') as provider:
            process_delivery(claim_delivery())
        delivery.refresh_from_db()
        self.assertEqual(delivery.status, 'blocked')
        provider.assert_not_called()

    def test_authorized_non_it_exporter_can_send_without_schedule_access(self):
        from .services import can_manage
        AccessGrant.objects.create(user=self.denied, workflow='tat_tracker', role='MANAGEMENT', group_configuration=self.group)
        delivery = queue_export(self.denied, 'tat_tracker', {**self.payload, 'preset':'tat'})
        self.assertFalse(can_manage(self.denied, workflow='tat_tracker'))
        self.assertEqual(delivery.requested_by_id, self.denied.pk)

    def test_scope_change_blocks_captured_export(self):
        AccessGrant.objects.create(user=self.denied, workflow='tat_tracker', role='MANAGEMENT', group_configuration=self.group, branch='Training')
        delivery = queue_export(self.denied, 'tat_tracker', {**self.payload, 'preset':'tat'})
        AccessGrant.objects.filter(user=self.denied).update(branch='Other training branch')
        with self.assertRaises(PermissionDenied):
            validate_export(delivery)

    def test_one_off_worker_freezes_and_submits_once(self):
        from .resend import SubmissionError
        delivery = self.queue(payload={**self.payload, 'filters':{}})
        snapshot = {'preset': 'pipeline', 'summary': {'Cases in scope': 1}, 'total_rows': 1,
                    'rows': [], 'xlsx_content': base64.b64encode(b'synthetic-workbook').decode(),
                    'run_at': '2026-10-05T09:00:00+03:00', 'applied_filters': {},
                    'charts': [{'title': 'Pipeline', 'labels': ['Visit'],
                                'datasets': [{'label': 'Cases', 'values': [1]}]}]}
        with patch('report_delivery.services.capture_report', return_value=snapshot), \
             patch('report_delivery.resend.send_email', side_effect=SubmissionError('network', retryable=True, ambiguous=True)) as provider:
            process_delivery(claim_delivery())
        delivery.refresh_from_db()
        original = delivery.payload
        self.assertEqual(delivery.status, 'retry')
        self.assertEqual(original['to'], [self.payload['email']])
        self.assertEqual(original['attachments'][2]['content_id'], 'jawabu-report-chart-1')
        from django.utils import timezone
        ReportDelivery.objects.filter(pk=delivery.pk).update(next_attempt_at=timezone.now())
        with patch('report_delivery.services.capture_report') as capture, \
             patch('report_delivery.rendering.build_payload') as render, \
             patch('report_delivery.resend.send_email', return_value='synthetic-provider') as provider:
            process_delivery(claim_delivery())
        capture.assert_not_called()
        render.assert_not_called()
        self.assertEqual(provider.call_args.args[0], original)
        delivery.refresh_from_db()
        self.assertEqual(delivery.payload, original)
        self.assertEqual(delivery.status, 'accepted')
        self.assertIsNone(claim_delivery())

    def test_signed_bounce_for_one_off_records_suppression(self):
        from django.utils import timezone
        delivery = self.queue()
        delivery.provider_id='synthetic-bounce'; delivery.status='accepted'; delivery.save()
        WebhookReceipt.objects.create(event_id='synthetic-event',provider_id=delivery.provider_id,event_type='email.bounced',occurred_at=timezone.now())
        reconcile_webhooks(delivery.provider_id)
        delivery.refresh_from_db()
        self.assertEqual(delivery.status,'bounced')
        with self.assertRaises(ValidationError):
            self.queue(payload={**self.payload,'client_request_id':str(uuid.uuid4())})

    def test_suppression_applies_to_new_destinations(self):
        delivery = self.queue()
        delivery.status = 'bounced'; delivery.save()
        with self.assertRaises(ValidationError):
            self.queue(payload={**self.payload, 'client_request_id': str(uuid.uuid4())})

    def test_status_is_owner_bound_and_disabled_delivery_is_explicit(self):
        delivery = self.queue()
        other = get_user_model().objects.create_user(username='synthetic-other', is_superuser=True)
        with patch('report_delivery.dispatch.wake_deliveries'):
            response = report_export_response(other, 'jawabu_portal', {'action':'status', 'delivery_id':str(delivery.pk)})
        self.assertEqual(response.status_code, 404)
        with override_settings(REPORT_EMAIL_DELIVERY_ENABLED=False):
            self.assertEqual(report_export_response(self.actor, 'jawabu_portal', self.payload).status_code, 400)

    def test_portal_attachment_is_the_existing_filtered_export(self):
        JawabuFarmerMaster.objects.create(customer_name='Synthetic matching case', group_configuration=self.group)
        JawabuFarmerMaster.objects.create(customer_name='Synthetic excluded case', group_configuration=self.group)
        delivery = self.queue(payload={**self.payload, 'filters': {'search': 'Synthetic matching'}})
        snapshot = capture_export(delivery)
        workbook = load_workbook(BytesIO(base64.b64decode(snapshot['xlsx_content'])))
        self.assertEqual(snapshot['total_rows'], 1)
        self.assertEqual(workbook['Data'].max_row, 2)
        payload = build_payload(snapshot, delivery.destination, delivery.configuration)
        self.assertEqual(payload['from'], 'JBL BOT <reports@example.invalid>')
        self.assertGreaterEqual(len(payload['attachments']), 2)
        self.assertEqual(payload['attachments'][1]['content_id'], 'jawabu-report-logo')
        self.assertIn('Detailed Excel attached', payload['html'])
        self.assertNotIn('Synthetic matching case', payload['html'])

    def test_native_exporters_keep_tat_and_complaint_workbook_bytes(self):
        for workflow, preset, module, exporter, summary in [
            ('tat_tracker', 'tat', 'core.services.tat_reporting', 'export_report_xlsx', 'report_summary'),
            ('complaint_cases', 'complaints', 'core.services.complaint_register', 'export_register_xlsx', 'complaint_report_summary')]:
            with self.subTest(workflow=workflow):
                delivery = self.queue(workflow, {**self.payload, 'client_request_id':str(uuid.uuid4()), 'preset':preset,
                                               'filters':{'branch':'Training', 'metric':'hb_response', 'metric_value':'late'}})
                with patch(module+'.'+exporter, return_value=(b'synthetic-native-xlsx', 4)) as export, patch(module+'.'+summary, return_value={'metrics':{'active':3,'stalled':1}}):
                    snapshot = capture_export(delivery)
                self.assertEqual(base64.b64decode(snapshot['xlsx_content']), b'synthetic-native-xlsx')
                self.assertEqual(snapshot['exported_rows'], 4)
                supplied = export.call_args.kwargs.get('filters', export.call_args.args[1] if len(export.call_args.args)>1 else {})
                self.assertEqual(supplied['metric_value'], 'late')
                if workflow == 'tat_tracker':
                    self.assertEqual(snapshot['summary']['Active cases'], 3)
                    self.assertEqual(snapshot['summary']['Stalled cases'], 1)

    def test_email_escape_dates_missing_values_and_payload_limit(self):
        from .rendering import ReportTooLarge
        snapshot = {'preset':'complaints', 'period':{'from':'2026-10-01','to':'2026-10-05'},
                    'summary':{'Open':0,'Median HB response (hours)':None}, 'charts':[], 'rows':[],
                    'total_rows':0, 'export_limit':None, 'xlsx_content':base64.b64encode(b'xlsx').decode(),
                    'applied_filters':{'search':'<script>unsafe</script>'}, 'run_at':'2026-10-05T09:00:00+03:00'}
        payload = build_payload(snapshot, self.payload['email'], {})
        self.assertIn('01-Oct-2026', payload['subject'])
        self.assertNotIn('<script>unsafe</script>', payload['html'])
        self.assertIn('&lt;script&gt;', payload['html'])
        self.assertIn('—', payload['html'])
        snapshot.update(preset='tat', period=None, applied_filters={'view':'current','date_from':'2026-10-01'})
        current = build_payload(snapshot, self.payload['email'], {})
        self.assertIn('From 01-Oct-2026', current['subject'])
        self.assertIn('Current workload at generation time', current['html'])
        self.assertNotIn('Completed-period', current['html'])
        with patch('report_delivery.rendering.MAX_PAYLOAD_BYTES', 1), self.assertRaises(ReportTooLarge):
            build_payload(snapshot, self.payload['email'], {})
