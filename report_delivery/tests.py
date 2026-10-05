"""Synthetic-only delivery, disclosure, period and webhook regressions."""
import base64
import hashlib
import hmac
import json
import uuid
from datetime import datetime, timedelta
from io import BytesIO
from unittest.mock import Mock, patch
from zoneinfo import ZoneInfo

from django.contrib.auth import get_user_model
from django.core.exceptions import PermissionDenied, ValidationError
from django.core.management import call_command
from django.test import TestCase, TransactionTestCase, RequestFactory, override_settings
from django.urls import reverse
from django.utils import timezone
from openpyxl import load_workbook

from core.models import AccessGrant, GroupSheetConfiguration, JawabuFarmerMaster
from .models import ApprovedRecipient, ReportSchedule, ReportDelivery, WebhookReceipt
from .resend import SubmissionError, send_email, verify_webhook
from .services import (can_manage, capture_report, claim_delivery, completed_period,
                       latest_occurrence, next_occurrence, process_delivery, purge_history,
                       queue_schedule, reconcile_webhooks, reserve_due, retry_delivery,
                       validate_recipient)

EAT = ZoneInfo('Africa/Nairobi')
NOW = datetime(2026, 10, 5, 9, tzinfo=EAT)


@override_settings(REPORT_EMAIL_DELIVERY_ENABLED=True, RESEND_API_KEY='synthetic-key', REPORT_EMAIL_FROM='Reports <reports@example.invalid>')
class DeliveryTests(TransactionTestCase):
    def setUp(self):
        clock = patch('django.utils.timezone.now', return_value=NOW)
        clock.start()
        self.addCleanup(clock.stop)
        self.actor = get_user_model().objects.create_user(username='synthetic-report-admin', is_staff=True, is_superuser=True)
        self.officer = get_user_model().objects.create_user(username='synthetic-report-officer')
        self.group = GroupSheetConfiguration.objects.create(group_id='synthetic-report-group')
        self.other_group = GroupSheetConfiguration.objects.create(group_id='synthetic-report-other-group')
        self.recipient = ApprovedRecipient.objects.create(email='approved@example.invalid', group_configuration=self.group, authorized_by=self.actor)
        self.schedule = ReportSchedule.objects.create(title='Synthetic report', preset='pipeline', group_configuration=self.group, authorized_by=self.actor, active=True, next_run_at=NOW - timedelta(days=3))
        self.schedule.recipients.add(self.recipient)

    def queue(self):
        return queue_schedule(self.schedule.pk, actor=self.actor, request_key=uuid.uuid4(), now=NOW)[0]

    def grant(self, *, branch='', product='', group=None):
        return AccessGrant.objects.create(user=self.officer, workflow='jawabu_portal', role='IT', branch=branch, product=product, group_configuration=group or self.group)

    def test_single_complete_grant_is_required(self):
        self.grant(branch='A', product='X')
        self.grant(branch='B', product='Y')
        self.schedule.branch, self.schedule.product = 'A', 'Y'
        self.assertFalse(can_manage(self.officer, self.schedule))
        self.schedule.product = 'X'
        self.assertTrue(can_manage(self.officer, self.schedule))
        self.schedule.branch = ''
        self.assertFalse(can_manage(self.officer, self.schedule))

    def test_officer_cannot_queue_or_view_scope(self):
        with self.assertRaises(PermissionDenied):
            queue_schedule(self.schedule.pk, actor=self.officer, request_key=uuid.uuid4())

    def test_recipient_scope_cannot_be_widened(self):
        self.recipient.branch = 'A'
        self.recipient.save()
        with self.assertRaises(ValidationError):
            validate_recipient(self.schedule, self.recipient)
        self.schedule.branch = 'A'
        validate_recipient(self.schedule, self.recipient)
        self.recipient.group_configuration = self.other_group
        with self.assertRaises(ValidationError):
            validate_recipient(self.schedule, self.recipient)

    def test_duplicate_manual_action_reuses_delivery(self):
        key = uuid.uuid4()
        first = queue_schedule(self.schedule.pk, actor=self.actor, request_key=key, now=NOW)
        second = queue_schedule(self.schedule.pk, actor=self.actor, request_key=key, now=NOW)
        self.assertEqual(first[0].pk, second[0].pk)
        self.assertEqual(ReportDelivery.objects.count(), 1)

    def test_each_recipient_has_separate_reservation(self):
        other = ApprovedRecipient.objects.create(email='second@example.invalid', group_configuration=self.group, authorized_by=self.actor)
        self.schedule.recipients.add(other)
        self.assertEqual(len(queue_schedule(self.schedule.pk, actor=self.actor, request_key=uuid.uuid4())), 2)

    def test_latest_period_only_and_due_reservation_idempotent(self):
        self.assertEqual(reserve_due(now=NOW), 1)
        self.assertEqual(reserve_due(now=NOW), 0)
        self.schedule.refresh_from_db()
        self.assertEqual(self.schedule.skipped_occurrences, 3)
        self.assertGreater(self.schedule.next_run_at, NOW)
        self.assertEqual(ReportDelivery.objects.get().configuration['period'], {'from': '2026-10-04', 'to': '2026-10-04'})

    def test_daily_weekly_monthly_quarterly_boundaries(self):
        examples = [
            ('daily', NOW, '2026-10-04', '2026-10-04'),
            ('weekly', NOW, '2026-09-28', '2026-10-04'),
            ('monthly', datetime(2026, 3, 1, 8, tzinfo=EAT), '2026-02-01', '2026-02-28'),
            ('quarterly', datetime(2026, 1, 1, 8, tzinfo=EAT), '2025-10-01', '2025-12-31'),
        ]
        for frequency, due, start, end in examples:
            self.schedule.frequency = frequency
            self.assertEqual(completed_period(self.schedule, due), {'from': start, 'to': end})
            self.assertGreater(next_occurrence(self.schedule, due), due)
        self.schedule.frequency = 'daily'
        early = datetime(2026, 10, 5, 4, tzinfo=ZoneInfo('UTC'))  # 07:00 Nairobi
        self.assertEqual(latest_occurrence(self.schedule, early).date().isoformat(), '2026-10-04')

    def test_leap_month(self):
        self.schedule.frequency = 'monthly'
        self.assertEqual(completed_period(self.schedule, datetime(2028, 3, 1, 8, tzinfo=EAT))['to'], '2028-02-29')

    def test_manual_month_report_before_dispatch_time_uses_last_completed_month(self):
        self.schedule.frequency = 'monthly'
        self.schedule.save()
        early = datetime(2026, 10, 1, 7, tzinfo=EAT)
        delivery = queue_schedule(self.schedule.pk, actor=self.actor, request_key=uuid.uuid4(), now=early)[0]
        self.assertEqual(delivery.configuration['period'], {'from': '2026-09-01', 'to': '2026-09-30'})

    def test_exclusive_lease_and_expired_recovery(self):
        self.queue()
        delivery = claim_delivery(now=NOW)
        self.assertIsNotNone(delivery)
        self.assertIsNone(claim_delivery(now=NOW))
        recovered = claim_delivery(now=NOW + timedelta(minutes=6))
        self.assertEqual(delivery.pk, recovered.pk)
        self.assertNotEqual(delivery.lease_token, recovered.lease_token)

    def test_revoked_authorizer_blocks_before_render_or_send(self):
        self.grant()
        self.schedule.authorized_by = self.officer
        self.schedule.save()
        self.queue()
        AccessGrant.objects.filter(user=self.officer).update(active=False)
        delivery = claim_delivery(now=NOW)
        with patch('report_delivery.resend.send_email') as sender, patch('report_delivery.services.capture_report') as capture:
            process_delivery(delivery, now=NOW)
        sender.assert_not_called()
        capture.assert_not_called()
        delivery.refresh_from_db()
        self.assertEqual(delivery.status, 'blocked')

    def test_changed_recipient_address_blocks_frozen_work(self):
        self.queue()
        ApprovedRecipient.objects.filter(pk=self.recipient.pk).update(email='changed@example.invalid')
        delivery = claim_delivery(now=NOW)
        with patch('report_delivery.resend.send_email') as sender:
            process_delivery(delivery, now=NOW)
        sender.assert_not_called()
        delivery.refresh_from_db()
        self.assertEqual(delivery.status, 'blocked')

    def frozen(self, **values):
        delivery = self.queue()
        ReportDelivery.objects.filter(pk=delivery.pk).update(payload={'to': ['approved@example.invalid'], 'subject': 'Synthetic frozen'}, **values)
        return claim_delivery(now=NOW)

    def test_acceptance_and_retry_preserve_exact_payload_and_key(self):
        delivery = self.frozen()
        with patch('report_delivery.resend.send_email', side_effect=SubmissionError('network', retryable=True, ambiguous=True)) as sender:
            process_delivery(delivery, now=NOW)
        original = sender.call_args
        delivery.refresh_from_db()
        self.assertEqual(delivery.status, 'retry')
        delivery = claim_delivery(now=NOW + timedelta(minutes=2))
        with patch('report_delivery.resend.send_email', return_value='synthetic-provider-id') as sender:
            process_delivery(delivery, now=NOW + timedelta(minutes=2))
        self.assertEqual(original, sender.call_args)
        delivery.refresh_from_db()
        self.assertEqual(delivery.status, 'accepted')
        self.assertEqual(delivery.attempts, 2)

    def test_expired_submission_never_resends(self):
        delivery = self.frozen(first_attempt_at=NOW - timedelta(days=1))
        with patch('report_delivery.resend.send_email') as sender:
            process_delivery(delivery, now=NOW)
        sender.assert_not_called()
        delivery.refresh_from_db()
        self.assertEqual(delivery.status, 'uncertain')
        with self.assertRaises(ValidationError):
            retry_delivery(delivery.pk, self.actor)

    def test_generation_failure_is_not_an_empty_report(self):
        self.schedule.skip_empty = True
        self.schedule.save()
        self.queue()
        delivery = claim_delivery(now=NOW)
        with patch('report_delivery.services.capture_report', side_effect=ValueError('Synthetic generation error')):
            process_delivery(delivery, now=NOW)
        delivery.refresh_from_db()
        self.assertEqual(delivery.status, 'failed')

    def test_group_capture_and_excel_are_frozen(self):
        case = JawabuFarmerMaster.objects.create(customer_name='Synthetic allowed', group_configuration=self.group)
        JawabuFarmerMaster.objects.create(customer_name='Synthetic excluded', group_configuration=self.other_group)
        snapshot = capture_report(self.schedule, self.queue().configuration)
        self.assertEqual(snapshot['total_rows'], 1)
        self.assertEqual(snapshot['rows'][0]['customer_name'], 'Synthetic allowed')
        case.customer_name = 'Changed after capture'
        case.save()
        from core.services.portal_reporting import curated_snapshot_xlsx
        workbook = load_workbook(BytesIO(curated_snapshot_xlsx(snapshot)))
        self.assertEqual(workbook['Data']['B2'].value, 'Synthetic allowed')

    def test_formula_like_names_are_strings_in_excel(self):
        JawabuFarmerMaster.objects.create(customer_name='=SYNTHETIC()', group_configuration=self.group)
        snapshot = capture_report(self.schedule, self.queue().configuration)
        from core.services.portal_reporting import curated_snapshot_xlsx
        workbook = load_workbook(BytesIO(curated_snapshot_xlsx(snapshot)))
        self.assertEqual(workbook['Data']['B2'].data_type, 's')

    def test_webhook_out_of_order_and_suppression(self):
        delivery = self.queue()
        ReportDelivery.objects.filter(pk=delivery.pk).update(provider_id='synthetic-id', status='accepted')
        WebhookReceipt.objects.create(event_id='delivered', provider_id='synthetic-id', event_type='email.delivered', occurred_at=NOW)
        WebhookReceipt.objects.create(event_id='late-delay', provider_id='synthetic-id', event_type='email.delivery_delayed', occurred_at=NOW + timedelta(seconds=5))
        reconcile_webhooks('synthetic-id')
        delivery.refresh_from_db()
        self.assertEqual(delivery.status, 'delivered')
        WebhookReceipt.objects.create(event_id='bounce', provider_id='synthetic-id', event_type='email.bounced', occurred_at=NOW - timedelta(minutes=1))
        reconcile_webhooks('synthetic-id')
        self.recipient.refresh_from_db()
        self.assertTrue(self.recipient.suppressed)
        with self.assertRaises(ValidationError):
            validate_recipient(self.schedule, self.recipient)

    def test_retention_erases_data_then_metadata(self):
        delivery = self.queue()
        ReportDelivery.objects.filter(pk=delivery.pk).update(payload={'synthetic': 'attachment'}, snapshot={'synthetic': 'data'}, created_at=NOW - timedelta(days=31))
        self.assertEqual(purge_history(now=NOW)[0], 1)
        delivery.refresh_from_db()
        self.assertEqual(delivery.payload, {})
        self.assertEqual(delivery.snapshot, {})
        ReportDelivery.objects.filter(pk=delivery.pk).update(created_at=NOW - timedelta(days=181))
        self.assertEqual(purge_history(now=NOW)[1], 1)

    def test_disabled_runner_does_not_send(self):
        from django.core.management.base import CommandError
        self.queue()
        with override_settings(REPORT_EMAIL_DELIVERY_ENABLED=False), self.assertRaises(CommandError), patch('requests.post') as sender:
            call_command('process_report_deliveries', apply=True)
        sender.assert_not_called()

    def test_admin_controls_scope_and_no_send_during_request(self):
        self.client.force_login(self.actor)
        url = reverse('admin:report_schedule_controls', args=[self.schedule.pk])
        self.assertEqual(self.client.get(url).status_code, 200)
        with patch('report_delivery.resend.send_email') as sender:
            response = self.client.post(url, {'action': 'send', 'request_key': str(uuid.uuid4())})
        self.assertEqual(response.status_code, 302)
        sender.assert_not_called()
        self.officer.is_staff = True
        self.officer.save()
        self.grant(branch='A')
        self.client.force_login(self.officer)
        self.assertEqual(self.client.get(url).status_code, 404)

    def test_admin_schedule_form_exposes_approved_recipients(self):
        self.client.force_login(self.actor)
        response = self.client.get(reverse('admin:report_delivery_reportschedule_add'))
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, 'name="recipients"')

    def test_safe_retry_and_permanent_provider_rejection(self):
        delivery = self.frozen()
        with patch('report_delivery.resend.send_email', side_effect=SubmissionError('provider_rejected_request')):
            process_delivery(delivery, now=NOW)
        delivery.refresh_from_db()
        self.assertEqual(delivery.status, 'failed')
        retry_delivery(delivery.pk, self.actor)
        delivery.refresh_from_db()
        self.assertEqual(delivery.status, 'queued')

    def test_schedule_revision_change_blocks_old_delivery(self):
        self.queue()
        ReportSchedule.objects.filter(pk=self.schedule.pk).update(revision=2)
        delivery = claim_delivery(now=NOW)
        with patch('report_delivery.resend.send_email') as sender:
            process_delivery(delivery, now=NOW)
        sender.assert_not_called()
        delivery.refresh_from_db()
        self.assertEqual(delivery.status, 'blocked')

    def test_skip_empty_is_opt_in(self):
        self.schedule.skip_empty = True
        self.schedule.save()
        self.queue()
        delivery = claim_delivery(now=NOW)
        with patch('report_delivery.services.capture_report', return_value={'total_rows': 0}), patch('report_delivery.resend.send_email') as sender:
            process_delivery(delivery, now=NOW)
        sender.assert_not_called()
        delivery.refresh_from_db()
        self.assertEqual(delivery.status, 'skipped')

    def test_worker_losing_lease_during_capture_does_not_send(self):
        self.queue()
        delivery = claim_delivery(now=NOW)
        snapshot = {'total_rows': 1}
        def steal(*args, **kwargs):
            ReportDelivery.objects.filter(pk=delivery.pk).update(lease_token=uuid.uuid4())
            return snapshot
        with patch('report_delivery.services.capture_report', side_effect=steal), patch('report_delivery.rendering.build_payload', return_value={'synthetic': 'payload'}), patch('report_delivery.resend.send_email') as sender:
            process_delivery(delivery, now=NOW)
        sender.assert_not_called()

    def test_deleted_approving_account_blocks_delivery(self):
        from core.services.user_hard_delete import governed_user_hard_delete
        delivery = self.queue()
        with governed_user_hard_delete():
            self.actor.delete()
        delivery = claim_delivery(now=NOW)
        with patch('report_delivery.resend.send_email') as sender:
            process_delivery(delivery, now=NOW)
        sender.assert_not_called()
        delivery.refresh_from_db()
        self.assertEqual(delivery.status, 'blocked')

    def test_report_payload_has_both_attachments_same_facts(self):
        from .rendering import build_payload
        from pypdf import PdfReader
        JawabuFarmerMaster.objects.create(customer_name='Synthetic report case', group_configuration=self.group)
        configuration = self.queue().configuration
        snapshot = capture_report(self.schedule, configuration)
        payload = build_payload(snapshot, self.recipient.email, configuration)
        self.assertEqual(payload['to'], [self.recipient.email])
        self.assertEqual(len(payload['attachments']), 2)
        pdf = PdfReader(BytesIO(base64.b64decode(payload['attachments'][0]['content'])))
        text = '\n'.join(page.extract_text() for page in pdf.pages)
        self.assertIn('1 matching cases', text)
        workbook = load_workbook(BytesIO(base64.b64decode(payload['attachments'][1]['content'])))
        self.assertEqual(workbook['Data'].max_row - 1, len(snapshot['rows']))
        self.assertIn('Cases in scope', payload['html'])

    def test_scheduled_empty_is_sent_by_default(self):
        self.queue()
        delivery = claim_delivery(now=NOW)
        with patch('report_delivery.services.capture_report', return_value={'total_rows': 0}), patch('report_delivery.rendering.build_payload', return_value={'synthetic': 'payload'}), patch('report_delivery.resend.send_email', return_value='synthetic-empty-id') as sender:
            process_delivery(delivery, now=NOW)
        sender.assert_called_once()
        delivery.refresh_from_db()
        self.assertEqual(delivery.status, 'accepted')

    def test_settings_trigger_queues_and_rejects_officers(self):
        from core.api.portal_views import portal_report_send_now
        request = RequestFactory().post('/api/portal/settings/reports/send/', data=json.dumps({'client_request_id': str(uuid.uuid4())}), content_type='application/json')
        request.portal_user = self.actor
        with patch('report_delivery.resend.send_email') as sender:
            response = portal_report_send_now.__wrapped__.__wrapped__.__wrapped__(request)
        self.assertEqual(response.status_code, 202)
        sender.assert_not_called()
        request.portal_user = self.officer
        self.assertEqual(portal_report_send_now.__wrapped__.__wrapped__.__wrapped__(request).status_code, 403)


class AdapterTests(TestCase):
    @override_settings(RESEND_WEBHOOK_SECRET='whsec_c3ludGhldGljLXNlY3JldA==')
    def test_signed_webhook_duplicate_and_tampering(self):
        body = json.dumps({'type': 'email.delivered', 'created_at': timezone.now().isoformat(), 'data': {'email_id': 'synthetic-provider'}}).encode()
        timestamp = str(int(timezone.now().timestamp()))
        signature = base64.b64encode(hmac.new(b'synthetic-secret', b'event.' + timestamp.encode() + b'.' + body, hashlib.sha256).digest()).decode()
        headers = {'HTTP_SVIX_ID': 'event', 'HTTP_SVIX_TIMESTAMP': timestamp, 'HTTP_SVIX_SIGNATURE': 'v1,' + signature}
        url = reverse('report_resend_webhook')
        for _ in range(2):
            self.assertEqual(self.client.post(url, data=body, content_type='application/json', **headers).status_code, 200)
        self.assertEqual(WebhookReceipt.objects.count(), 1)
        self.assertEqual(self.client.post(url, data=body + b' ', content_type='application/json', **headers).status_code, 400)
        headers['HTTP_SVIX_TIMESTAMP'] = '1'
        self.assertEqual(self.client.post(url, data=body, content_type='application/json', **headers).status_code, 400)

    @override_settings(RESEND_WEBHOOK_SECRET='whsec_plJ3nmyCDGBKInavdOK15jsl')
    def test_official_svix_protocol_vector(self):
        headers = {'svix-id': 'msg_loFOjxBNrRLzqYUf', 'svix-timestamp': '1731705121', 'svix-signature': 'v1,rAvfW3dJ/X/qxhsaXPOyyCGmRKsaKWcsNccKXlIktD0='}
        self.assertEqual(verify_webhook(b'{"event_type":"ping","data":{"success":true}}', headers, now=datetime.fromtimestamp(1731705121, ZoneInfo('UTC')))[0], headers['svix-id'])

    @override_settings(REPORT_EMAIL_DELIVERY_ENABLED=True, RESEND_API_KEY='synthetic', REPORT_EMAIL_FROM='reports@example.invalid')
    def test_provider_http_idempotency_and_errors(self):
        response = Mock(status_code=200)
        response.json.return_value = {'id': 'synthetic-id'}
        with patch('report_delivery.resend.requests.post', return_value=response) as sender:
            self.assertEqual(send_email({'to': ['approved@example.invalid']}, key='synthetic-key'), 'synthetic-id')
        self.assertEqual(sender.call_args.kwargs['headers']['Idempotency-Key'], 'synthetic-key')
        response.status_code = 429
        response.headers = {'Retry-After': '300'}
        with patch('report_delivery.resend.requests.post', return_value=response), self.assertRaises(SubmissionError) as error:
            send_email({}, key='synthetic-key')
        self.assertTrue(error.exception.retryable)
        self.assertEqual(error.exception.retry_after, 300)
