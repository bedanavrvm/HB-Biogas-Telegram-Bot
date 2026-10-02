"""Synthetic, local-only complaint reporting/timing regression cases."""
from datetime import datetime, timedelta
from io import BytesIO
from types import SimpleNamespace
from zoneinfo import ZoneInfo

from django.test import TestCase
from django.contrib.auth import get_user_model
from openpyxl import load_workbook

from core.models import RawMessage, ProcessedMessage, ParsedMessage, ComplaintCaseControl, CaseUpdate
from core.services.complaint_timing import case_timing, hydrate_timing, timing_summary, period_bucket
from core.services.complaint_register import (
    complaint_report_page, complaint_report_summary, export_register_xlsx,
    encode_export_filters, decode_export_filters, validate_report_filters,
)
from core.services.complaint_cases import ComplaintCaseError
from core.services.compliance_audit import record_event


class ComplaintTimingTests(TestCase):
    now = datetime(2026, 10, 2, 12, tzinfo=ZoneInfo('Africa/Nairobi'))

    def make_case(self, *, hours=48, closed=False, duration=24, target=24):
        index = ParsedMessage.objects.count() + 1
        raw = RawMessage.objects.create(telegram_message_id=f'timing-{index}', content='Synthetic training fixture')
        processed = ProcessedMessage.objects.create(message_hash=f'timing-{index}', raw_message=raw)
        start = self.now - timedelta(hours=hours)
        case = ParsedMessage.objects.create(
            processed_message=processed, message_id=f'timing-{index}', group_id='-100-timing',
            timestamp=start, sender='Training officer', raw_message='', customer_name=f'Training {index}',
            complaint_category='Training issue', complaint_status='Closed' if closed else 'Open',
            date_resolved=start + timedelta(hours=duration) if closed else None,
        )
        ComplaintCaseControl.objects.create(parsed_message=case, reference_number=f'CMP-TIMING-{index}',
            sla_target_hours=target, sla_started_at=start, sla_due_at=start + timedelta(hours=target))
        return case

    def update(self, case, hour, old, new, *, source='mini_app', affiliation=None):
        update = CaseUpdate.objects.create(parsed_message=case, group_id=case.group_id,
            old_status=old, new_status=new, source=source, raw_update_text='', updated_by='Training actor')
        CaseUpdate.objects.filter(pk=update.pk).update(created_at=case.timestamp + timedelta(hours=hour))
        if affiliation is not None:
            record_event(workflow='complaint_cases', action='case.updated', subject_type='complaint',
                subject_id=str(case.pk), source_model='CaseUpdate', source_event_id=str(update.pk),
                deduplication_key=f'timing-{update.pk}', metadata={'actor_affiliation': affiliation})
        return update

    def test_closed_timer_freezes_and_exact_deadline_is_on_time(self):
        case = self.make_case(closed=True)
        early = case_timing(case.complaint_control, case, now=self.now)
        later = case_timing(case.complaint_control, case, now=self.now + timedelta(days=60))
        self.assertEqual(early, later)
        self.assertEqual(early['resolution_hours'], 24)
        self.assertIs(early['on_time'], True)
        case.date_resolved += timedelta(seconds=1)
        self.assertIs(case_timing(case.complaint_control, case, now=self.now)['on_time'], False)

    def test_comments_and_reopening_keep_original_clock_and_history(self):
        case = self.make_case(hours=100)
        self.update(case, 10, 'Open', 'Open', source='mini_app_comment', affiliation='JBL')
        self.update(case, 12, 'Open', 'Open', source='mini_app_comment', affiliation='HB')
        self.update(case, 20, 'Open', 'Closed', affiliation='HB')
        self.update(case, 30, 'Closed', 'Reopened')
        self.update(case, 70, 'Reopened', 'Closed', affiliation='HB')
        self.update(case, 80, 'Closed', 'Reopened')
        case.complaint_status = 'Reopened'
        hydrate_timing([case], now=self.now)
        facts = case.complaint_timing
        self.assertEqual(facts['elapsed_seconds'], 100 * 3600)
        self.assertIsNone(facts['resolution_hours'])
        self.assertEqual(facts['first_resolution_hours'], 20)
        self.assertEqual(facts['hb_response_hours'], 12)
        summary = timing_summary([case], now=self.now, granularity='year')
        self.assertEqual(sum(row['resolved'] for row in summary['activity']), 2)
        self.assertEqual(sum(row['count'] for row in summary['reopenings']), 1)

    def test_unknown_roles_never_invent_hb_response(self):
        case = self.make_case(closed=True)
        self.update(case, 24, 'Open', 'Closed')
        hydrate_timing([case], now=self.now)
        self.assertIsNone(case.complaint_timing['hb_response_hours'])

    def test_history_queries_are_batched_and_missing_closure_never_grows(self):
        cases = [self.make_case(closed=True) for _ in range(10)]
        self.update(cases[0], 24, 'Open', 'Closed')
        for case in cases:
            # Cache the already-loaded controls; timing must not cause N+1 reads.
            case.complaint_control
        with self.assertNumQueries(2):
            hydrate_timing(cases, now=self.now)
        broken = cases[0]
        broken.date_resolved = None
        first = case_timing(broken.complaint_control, broken, now=self.now)
        later = case_timing(broken.complaint_control, broken, now=self.now + timedelta(days=30))
        self.assertEqual(first, later)
        self.assertIsNone(first['elapsed_seconds'])

    def test_medians_age_edges_and_invalid_timing(self):
        cases = [self.make_case(hours=240, closed=True, duration=value) for value in [2, 4, 200]]
        cases += [self.make_case(hours=days * 24) for days in [2, 3, 7, 8, 14, 15]]
        bad = self.make_case(closed=True, duration=-1)
        cases.append(bad)
        hydrate_timing(cases, now=self.now)
        summary = timing_summary(cases, now=self.now)
        self.assertEqual(summary['timing']['median_resolution_hours'], 4)
        self.assertEqual(summary['timing']['resolution_excluded'], 1)
        self.assertEqual([row['count'] for row in summary['open_age']], [1, 2, 2, 1])
        self.assertEqual(summary['timing']['on_time_percent'], 66.7)
        control = SimpleNamespace(sla_started_at=self.now, sla_target_hours=None)
        self.assertIsNone(case_timing(control, SimpleNamespace(complaint_status='Closed',
            date_resolved=self.now, timestamp=self.now, created_at=self.now), now=self.now)['on_time'])

    def test_nairobi_boundaries_and_resolution_date_cohort(self):
        self.assertEqual(period_bucket(datetime(2026, 3, 31, 22, tzinfo=ZoneInfo('UTC')), 'month'), '2026-04')
        case = self.make_case(hours=48, closed=True, duration=36)
        filters = {'date_basis': 'resolved', 'metric': 'resolution', 'date_from': '2026-10-02', 'date_to': '2026-10-02'}
        table = complaint_report_page(filters=filters)
        self.assertEqual(table['count'], 1)
        self.assertEqual(table['results'][0]['resolution_hours'], 36)
        self.assertEqual(complaint_report_page(filters={**filters, 'date_basis': 'reported', 'metric': ''})['count'], 0)
        summary = complaint_report_summary(filters={'date_from': '2026-10-02', 'date_to': '2026-10-02'})
        self.assertEqual(summary['timing']['resolution_count'], 1)
        self.assertEqual(sum(row['received'] for row in summary['activity']), 0)

    def test_filtered_export_matches_chart_selection_and_keeps_search_private(self):
        self.make_case(hours=96)
        self.make_case(hours=24)
        filters = {'metric': 'open_age', 'metric_value': '3_7', 'date_basis': 'reported'}
        self.assertEqual(complaint_report_page(filters=filters)['count'], 1)
        actor = get_user_model().objects.create_user(username='training-export')
        workbook, count = export_register_xlsx(actor=actor, request_id='training-export', filters=filters)
        self.assertEqual(count, 1)
        sheet = load_workbook(BytesIO(workbook)).active
        self.assertEqual(sheet.max_row, 2)
        self.assertEqual(sheet.cell(1, 22).value, 'Resolution Time (hours)')
        private = {'search': 'Synthetic name', **filters}
        encoded = encode_export_filters(private)
        self.assertNotIn('Synthetic name', encoded)
        self.assertEqual(decode_export_filters(encoded), private)
        for invalid in ({'metric': 'bad'}, {'metric': 'open_age', 'metric_value': 'bad'},
                        {'date_from': '2026-10-03', 'date_to': '2026-10-02'},
                        {'metric': 'resolution', 'date_basis': 'reported'}):
            with self.assertRaises(ComplaintCaseError):
                validate_report_filters(invalid)
