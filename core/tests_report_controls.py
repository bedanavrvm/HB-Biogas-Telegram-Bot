"""Synthetic regressions for report cohorts, export scope and heatmap counts."""
from datetime import timedelta
from io import BytesIO

from django.contrib.auth import get_user_model
from django.test import SimpleTestCase, TestCase
from django.utils import timezone
from openpyxl import load_workbook

from core.models import AccessGrant, GroupSheetConfiguration, TatTrackerCase

from core.services.tat_reporting import _filters, _heatmap_payload, _heatmap_dimension_labels, _chart_payload
from core.services.tat_reporting import export_report_xlsx, report_cases, report_summary


class ReportHeatmapContractTests(SimpleTestCase):
    def test_chart_defaults_match_the_measurement(self):
        def chart(key, **extras):
            return _chart_payload(key, 'Synthetic', '', '', [], [], applied_filters=[], extras=extras)
        self.assertEqual(chart('trend')['default_type'], 'line')
        self.assertEqual(chart('stage_target')['default_type'], 'bar')
        self.assertEqual(chart('explorer', metric='sla_state')['default_type'], 'stacked_bar')
        self.assertIn('doughnut', chart('backlog_age')['allowed_types'])

    def test_mixed_paths_keep_disbursement_last(self):
        samples = [
            {'stage_key': 'disbursement', 'stage': 'Finance disbursement', '_stage_order': 3},
            {'stage_key': 'tat_held', 'stage': 'HOCC held', '_stage_order': 8},
            {'stage_key': 'ca_analysis_sent', 'stage': 'Credit analysis sent', '_stage_order': 2},
        ]
        self.assertEqual(_heatmap_dimension_labels(samples, 'stage', sample=True),
                         ['Credit analysis sent', 'HOCC held', 'Finance disbursement'])

    def test_case_count_is_distinct_and_missing_targets_are_excluded(self):
        samples = [
            {'_case_pk': 'a', 'stage': 'Credit', 'branch': 'Training', 'sla_state': 'within_target'},
            {'_case_pk': 'a', 'stage': 'Credit', 'branch': 'Training', 'sla_state': 'overdue'},
            {'_case_pk': 'b', 'stage': 'Credit', 'branch': 'Training', 'sla_state': 'target_unavailable'},
        ]
        filters = _filters({'view': 'performance'})
        cell = _heatmap_payload([], samples, filters)['cells'][0]
        self.assertEqual(cell['value'], 50)
        self.assertEqual(cell['sample_count'], 2)
        self.assertEqual(cell['case_count'], 1)
        self.assertEqual(cell['excluded_count'], 1)

    def test_empty_heatmap_measurement_stays_unavailable(self):
        filters = _filters({'view': 'performance'})
        cell = _heatmap_payload([], [{'_case_pk': 'a', 'stage': 'Credit',
                                    'branch': 'Training', 'sla_state': 'target_unavailable'}], filters)['cells'][0]
        self.assertIsNone(cell['value'])
        self.assertEqual(cell['case_count'], 0)


class ReportScopeIntegrationTests(TestCase):
    def setUp(self):
        self.actor = get_user_model().objects.create_user(username='report-controls-synthetic')
        self.config = GroupSheetConfiguration.objects.create(group_id='-100-report-controls',
            display_name='Synthetic report', sheet_id='synthetic', sheet_name='Training',
            workflow={'type': 'tat_tracker', 'products': ['business'], 'branches': ['Training']})
        AccessGrant.objects.create(user=self.actor, workflow='tat_tracker', role='IT',
            group_configuration=self.config, branch='Training', product='business')

    def case(self, reference, *, branch='Training', completed=True):
        now = timezone.now()
        return TatTrackerCase.objects.create(group_id=self.config.group_id, case_id=reference,
            product_key='business', product_label='Business', client_name='Synthetic customer',
            branch=branch, status='Active', stage_values={
                'created': (now-timedelta(minutes=100)).isoformat(),
                **({'mpesa_to_admin': (now-timedelta(minutes=70)).isoformat()} if completed else {})},
            stage_target_snapshots={'mpesa_to_admin': {'target_minutes': '60'}})

    def test_all_export_ignores_dates_and_optional_filters_but_not_grants(self):
        old = self.case('TAT-SYNTHETIC-OLD', completed=False)
        old_time = timezone.now() - timedelta(days=450)
        TatTrackerCase.objects.filter(pk=old.pk).update(created_at=old_time)
        self.case('TAT-SYNTHETIC-NOW', completed=False)
        self.case('TAT-OUTSIDE-GRANT', branch='Other', completed=False)
        content, count = export_report_xlsx(self.actor, {'view': 'performance',
            'export_scope': 'all', 'branch': 'Other', 'search': 'not-a-match',
            'date_from': '2026-10-01', 'date_to': '2026-10-01'}, request_id='synthetic-all')
        self.assertEqual(count, 2)
        workbook = load_workbook(BytesIO(content))
        self.assertEqual({workbook['TAT Report'].cell(row, 1).value for row in (2, 3)},
                         {'TAT-SYNTHETIC-OLD', 'TAT-SYNTHETIC-NOW'})

    def test_heatmap_case_count_matches_drilldown_and_export(self):
        self.case('TAT-SYNTHETIC-A')
        self.case('TAT-SYNTHETIC-B')
        self.case('TAT-SYNTHETIC-NOT-COMPLETED', completed=False)
        self.case('TAT-SYNTHETIC-OUTSIDE', branch='Other')
        filters = {'view': 'performance', 'heatmap_metric': 'sla_met'}
        heatmap = report_summary(self.actor, filters)['heatmap']
        cell = next(cell for cell in heatmap['cells'] if cell['case_count'])
        selection = {**filters, 'heat_row': cell['row'], 'heat_column': cell['column']}
        result = report_cases(self.actor, selection)
        self.assertEqual(result['count'], cell['case_count'])
        self.assertEqual({row['case_id'] for row in result['results']},
                         {'TAT-SYNTHETIC-A', 'TAT-SYNTHETIC-B'})
        self.assertNotIn('_case_pk', result['results'][0])
        _, count = export_report_xlsx(self.actor, selection, request_id='synthetic-cell')
        self.assertEqual(count, cell['case_count'])

    def test_no_completed_actions_does_not_fall_back_to_current_workload(self):
        self.case('TAT-SYNTHETIC-OPEN', completed=False)
        performance = report_summary(self.actor, {'view': 'performance'})
        self.assertEqual(performance['breakdown_basis'], 'completed_stage_actions')
        self.assertEqual(performance['by_stage'], [])
        self.assertEqual(performance['by_role'], [])
        self.assertTrue(report_summary(self.actor, {'view': 'current'})['by_stage'])

    def test_incomplete_or_unknown_cell_cannot_widen_the_cohort(self):
        self.case('TAT-SYNTHETIC-A')
        with self.assertRaisesMessage(ValueError, 'complete heatmap cell'):
            report_cases(self.actor, {'heat_row': 'Training'})
        self.assertEqual(report_cases(self.actor, {'heat_row': 'unknown', 'heat_column': 'Training'})['count'], 0)
        with self.assertRaisesMessage(ValueError, 'filtered records or all records'):
            export_report_xlsx(self.actor, {'export_scope': 'unknown'})

    def test_chart_selection_is_scoped_paged_and_exported(self):
        self.case('TAT-CHART-DONE')
        self.case('TAT-CHART-OPEN', completed=False)
        self.case('TAT-CHART-OUTSIDE', branch='Other')
        base = {'view': 'performance', 'chart_dimension': 'branch', 'chart_metric': 'workload'}
        chart = report_summary(self.actor, base)['charts']['explorer']
        self.assertTrue(chart['drilldown']['available'])
        selection = {**base, 'drill_chart': 'explorer', 'drill_series': 'count', 'drill_bucket': 'Training'}
        result = report_cases(self.actor, selection)
        self.assertEqual([row['case_id'] for row in result['results']], ['TAT-CHART-DONE'])
        self.assertEqual(export_report_xlsx(self.actor, selection, request_id='training-chart')[1], 1)
        self.assertEqual(report_cases(self.actor, {**selection, 'drill_bucket': 'Other'})['count'], 0)

    def test_unknown_and_incomplete_chart_selections_are_rejected(self):
        self.case('TAT-CHART-OPEN', completed=False)
        for selection in [{'drill_chart': 'explorer'},
                          {'drill_chart': 'customer', 'drill_series': 'count', 'drill_bucket': 'Training'}]:
            with self.assertRaises(ValueError):
                report_cases(self.actor, selection)

    def test_historical_workload_cannot_masquerade_as_current_cases(self):
        chart = report_summary(self.actor, {})['charts']['trend']
        self.assertFalse(chart['drilldown']['available'])
        with self.assertRaisesMessage(ValueError, 'historical total'):
            report_cases(self.actor, {'drill_chart': 'trend', 'drill_series': 'active', 'drill_bucket': '2026-10-01'})

    def test_chart_metrics_select_their_denominator_and_keep_summary_unchanged(self):
        self.case('TAT-CHART-A')
        self.case('TAT-CHART-B')
        self.case('TAT-CHART-OUTSIDE', branch='Other')
        for metric in ['workload', 'sla_state', 'duration', 'target_usage', 'sla_met', 'correction_rate']:
            base = {'view':'performance', 'chart_dimension':'branch', 'chart_metric':metric}
            chart = report_summary(self.actor, base)['charts']['explorer']
            for series in chart['series']:
                selection = {**base, 'drill_chart':'explorer', 'drill_series':series['key'], 'drill_bucket':'Training'}
                rows = report_cases(self.actor, selection)['results']
                if metric == 'sla_state':
                    expected = 2 if series['values'][0] else 0
                    self.assertEqual(len(rows), expected)
                else:
                    # Percentiles, correction rates and percentages represent
                    # the entire measured cohort, not only numerator cases.
                    self.assertEqual(len(rows), 2)
                self.assertEqual(report_summary(self.actor, selection)['charts']['explorer'], chart)

    def test_time_and_target_points_retain_scopes_and_distinct_cases(self):
        self.case('TAT-CHART-A')
        self.case('TAT-CHART-B')
        self.case('TAT-CHART-OUTSIDE', branch='Other')
        base = {'view':'performance', 'stage':'mpesa_to_admin', 'granularity':'day'}
        charts = report_summary(self.actor, base)['charts']
        for key in ['trend', 'sla_compliance', 'tat_percentiles', 'stage_target']:
            chart = charts[key]
            for series in chart['series']:
                selection = {**base, 'drill_chart':key, 'drill_series':series['key'], 'drill_bucket':chart['labels'][0]}
                result = report_cases(self.actor, selection)
                self.assertEqual({row['case_id'] for row in result['results']}, {'TAT-CHART-A', 'TAT-CHART-B'})
                self.assertEqual(export_report_xlsx(self.actor, selection)[1], 2)
        with self.assertRaisesMessage(ValueError, 'supported chart series'):
            report_cases(self.actor, {**base, 'drill_chart':'trend', 'drill_series':'not-a-series', 'drill_bucket':'unknown'})

    def test_backlog_points_match_the_case_count_in_the_chart(self):
        self.case('TAT-BACKLOG-A', completed=False)
        self.case('TAT-BACKLOG-B', completed=False)
        self.case('TAT-BACKLOG-OUTSIDE', completed=False, branch='Other')
        chart = report_summary(self.actor, {'view': 'current'})['charts']['backlog_age']
        for index, bucket in enumerate(chart['labels']):
            result = report_cases(self.actor, {'view': 'current', 'drill_chart': 'backlog_age',
                'drill_series': 'cases', 'drill_bucket': bucket})
            self.assertEqual(result['count'], chart['series'][0]['values'][index])
