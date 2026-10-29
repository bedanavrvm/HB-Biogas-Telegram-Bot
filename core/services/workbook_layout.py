"""Preserve printed workbook geometry when a generated table changes size."""
from copy import copy
import textwrap

from openpyxl.utils import get_column_letter
from openpyxl.worksheet.print_settings import PrintArea


def shift_rows(worksheet, start: int, delta: int) -> None:
    """Move the layout metadata that openpyxl insert/delete_rows leaves behind."""
    if not delta:
        return

    def row_after(row):
        if row < start:
            return row
        if delta < 0 and row < start - delta:
            return None
        return row + delta

    def range_after(area):
        if area.max_row < start:
            return area
        if delta > 0:
            area.min_row += delta if area.min_row >= start else 0
            area.max_row += delta
        else:
            last_deleted = start - delta - 1
            if area.min_row >= start and area.max_row <= last_deleted:
                return None
            area.min_row = row_after(area.min_row) or start
            area.max_row = row_after(area.max_row) or start - 1
        return area

    merges = [copy(area) for area in worksheet.merged_cells.ranges if area.max_row >= start]
    for area in merges:
        worksheet.unmerge_cells(str(area))
    dimensions = {row: copy(dimension) for row, dimension in worksheet.row_dimensions.items() if row >= start}
    for row in dimensions:
        del worksheet.row_dimensions[row]
    print_areas = [copy(area) for area in PrintArea.from_string(str(worksheet.print_area or '')).ranges]
    if delta > 0:
        worksheet.insert_rows(start, delta)
    else:
        worksheet.delete_rows(start, -delta)
    for area in merges:
        moved = range_after(area)
        if moved is not None:
            worksheet.merge_cells(str(moved))
    for row, dimension in dimensions.items():
        moved = row_after(row)
        if moved is not None:
            dimension.index = moved
            worksheet.row_dimensions[moved] = dimension
    if print_areas:
        worksheet.print_area = [str(moved) for area in print_areas if (moved := range_after(area)) is not None]
    for image in getattr(worksheet, '_images', ()):
        for marker_name in ('_from', 'to'):
            marker = getattr(getattr(image, 'anchor', None), marker_name, None)
            if marker is not None:
                marker.row = (row_after(marker.row + 1) or start) - 1
    for page_break in worksheet.row_breaks.brk:
        page_break.id = row_after(page_break.id) or start


def configure_table_printing(worksheet, *, header_row: int, first_data_row: int, reference: str) -> None:
    """Keep all rows printable at one page wide, with readable continuation pages."""
    if not worksheet.print_area:
        worksheet.print_area = worksheet.calculate_dimension()
    worksheet.print_title_rows = f'{header_row}:{first_data_row - 1}'
    worksheet.page_setup.fitToWidth = 1
    worksheet.page_setup.fitToHeight = 0
    worksheet.page_setup.scale = None
    worksheet.sheet_properties.pageSetUpPr.fitToPage = True
    if not worksheet.page_setup.paperSize:
        worksheet.page_setup.paperSize = worksheet.PAPERSIZE_A4
    if not worksheet.oddHeader.right.text:
        worksheet.oddHeader.right.text = reference
    if not worksheet.oddFooter.right.text:
        worksheet.oddFooter.right.text = 'Page &P of &N'


def fit_wrapped_row_height(worksheet, row: int, columns) -> None:
    """Allow complete wrapped values, using each column's width and font size."""
    height = float(worksheet.row_dimensions[row].height or 18)
    for column in columns:
        cell = worksheet.cell(row, column)
        font_size = float(cell.font.sz or 11)
        width = float(worksheet.column_dimensions[get_column_letter(column)].width or 13)
        characters = max(1, int((width - 2) * 11 / font_size))
        lines = sum(max(1, len(textwrap.wrap(line, width=characters, break_on_hyphens=False)))
                    for line in str(cell.value or '').split('\n'))
        height = max(height, (font_size + 4) * lines + 4)
    worksheet.row_dimensions[row].height = min(409.5, height)
