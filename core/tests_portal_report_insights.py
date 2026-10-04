"""Synthetic Portal insight and drill/export consistency regressions."""
from datetime import datetime, timedelta
from decimal import Decimal
from io import BytesIO
from unittest.mock import patch
from zoneinfo import ZoneInfo

from django.contrib.auth import get_user_model
from django.db import connection
from django.test import TestCase
from django.test.utils import CaptureQueriesContext
from django.utils import timezone
from openpyxl import load_workbook

from core.models import AccessGrant, GroupSheetConfiguration, JawabuFarmerMaster, JawabuPipelineEvent, Product
from core.services.portal_reporting import run_curated_report, export_curated_report, PortalReportingError
from core.services.jawabu_case360 import calculate_case_tat


class PortalInsightTests(TestCase):
    def setUp(self):
        self.user = get_user_model().objects.create_user(username='synthetic-insight-superuser', is_superuser=True)
        self.now = timezone.now()
        self.day = timezone.localtime(self.now, ZoneInfo('Africa/Nairobi')).date()

    def case(self, **values):
        return JawabuFarmerMaster.objects.create(**{'customer_name': 'Synthetic reporting case', **values})

    def run_report(self, preset='pipeline', filters=None, user=None, access=None, page=1):
        return run_curated_report(preset=preset, filters=filters or {}, user=user or self.user, access=access, page=page)

    def chart(self, report, key):
        return next(c for c in report['charts'] if c['id'] == key)

    def event(self, record, action, at, **values):
        return JawabuPipelineEvent.objects.create(farmer=record, action=action, occurred_at=at, **values)

    def test_milestones_use_own_dates_and_nairobi_boundary(self):
        at = datetime(2026, 9, 30, 22, tzinfo=ZoneInfo('UTC'))  # 1 October Nairobi
        self.case(jbl_visit_date=at.date(), credit_decided_at=at, credit_decision='Approved')
        f = {'date_mode': 'custom', 'from': '2026-10-01', 'to': '2026-10-01', 'granularity': 'day'}
        report = self.run_report('outcomes', f)
        self.assertEqual(report['summary']['Visits'], 0)
        self.assertEqual(report['summary']['Credit decisions'], 1)
        self.assertEqual(report['summary']['Cases without timing'], 1)
        self.assertEqual(self.chart(report, 'duration')['unit'], 'hours')
        self.assertEqual(self.chart(report, 'activity')['bucket_keys'], ['2026-10-01'])

    def test_decimal_finance_selection_matches_table_and_export(self):
        first = self.case(branch='Synthetic A', requisition_date=self.day, invoice_amount=Decimal('12000.51'), payment=Decimal('3000.10'))
        self.case(branch='Synthetic B', invoice_date=self.day, invoice_amount=Decimal('0.09'))
        report = self.run_report('finance')
        self.assertEqual(report['summary']['Invoice amount'], '12000.60')
        self.assertEqual(self.chart(report, 'finance_branch')['datasets'][0]['values'], ['12000.51', '0.09'])
        f = {'chart_key': 'finance_branch', 'bucket_key': 'Synthetic A', 'series_key': 'Invoice value'}
        selected = self.run_report('finance', f)
        self.assertEqual(selected['total_rows'], 1)
        self.assertEqual(selected['rows'][0]['record_id'], str(first.pk))
        workbook = load_workbook(BytesIO(export_curated_report(preset='finance', filters=f, user=self.user, access=None)))
        self.assertEqual(workbook['Data'].max_row, 2)
        self.assertEqual(workbook['Data']['A2'].value, selected['rows'][0]['case_id'])

    def test_backlog_missing_timing_is_not_zero_and_legacy_stage_is_inferred(self):
        record = self.case()
        report = self.run_report()
        self.assertEqual(self.chart(report, 'stages')['bucket_keys'], ['jbl_visit'])
        self.assertEqual(report['rows'][0]['workflow_state'], 'jbl_visit')
        self.assertEqual(self.chart(report, 'age')['bucket_keys'], ['unavailable'])
        selected = self.run_report(filters={'stage': 'jbl_visit', 'chart_key': 'age', 'bucket_key': 'unavailable', 'series_key': 'Cases'})
        self.assertEqual(selected['rows'][0]['record_id'], str(record.pk))

    def test_duration_percentiles_targets_and_prefetched_calculator_parity(self):
        key = 'application_imported_to_jbl_visit_completed'
        for hours in (2, 4):
            record = self.case(jbl_visit_date=self.day)
            self.event(record, 'application_imported', self.now - timedelta(hours=hours), metadata={'tat_target_snapshot': {'stage_key': key, 'target_minutes': 180}})
            self.event(record, 'jbl_visit_completed', self.now)
            events = list(record.pipeline_events.order_by('occurred_at', 'created_at'))
            ordinary = calculate_case_tat(record, now=self.now)
            prefetched = calculate_case_tat(record, now=self.now, events=events, hb_events=[], targets={}, holidays=set())
            self.assertEqual(ordinary['stages'], prefetched['stages'])
        report = self.run_report('outcomes', {'date_mode': 'all'})
        chart = self.chart(report, 'duration')
        self.assertEqual(Decimal(chart['datasets'][0]['values'][0]), Decimal('3'))
        self.assertEqual(Decimal(chart['datasets'][1]['values'][0]), Decimal('3.8'))
        self.assertEqual(chart['sample_counts'][key], 2)
        sla = self.chart(report, 'sla')
        self.assertEqual({d['key']: sum(map(Decimal, d['values'])) for d in sla['datasets']}, {'On time': 1, 'Late': 1})

    def test_complete_grant_tuple_and_revocation_bound_every_projection(self):
        from core.services.telegram_identity import user_access
        user = get_user_model().objects.create_user(username='synthetic-report-it')
        group = GroupSheetConfiguration.objects.create(group_id='synthetic-report-group', workflow={'type': 'jawabu'})
        other = GroupSheetConfiguration.objects.create(group_id='synthetic-other-group', workflow={'type': 'jawabu'})
        product = Product.objects.create(code='SYNTHETIC-REPORT', name='Synthetic report product')
        wanted = self.case(branch='Synthetic A', product=product, group_configuration=group)
        self.case(branch='Synthetic A', product=product, group_configuration=other)
        self.case(branch='Synthetic B', product=product, group_configuration=group)
        grant = AccessGrant.objects.create(user=user, workflow='jawabu_portal', role='IT', branch='Synthetic A', product=product.code, group_configuration=group)
        access = user_access(user, 'jawabu_portal')
        report = self.run_report(user=user, access=access)
        self.assertEqual([r['record_id'] for r in report['rows']], [str(wanted.pk)])
        self.assertEqual(report['filter_options']['branches'], ['Synthetic A'])
        AccessGrant.objects.filter(pk=grant.pk).update(active=False)
        self.assertEqual(self.run_report(user=user, access=access)['total_rows'], 0)

    def test_table_cap_does_not_cap_chart_aggregates_or_filter_choices(self):
        for _ in range(3):
            self.case(branch='Synthetic A')
        with patch('core.services.portal_reporting.MAX_TABLE_ROWS', 1), patch('core.services.portal_reporting.PAGE_SIZE', 1):
            report = self.run_report()
        self.assertEqual(len(report['rows']), 1)
        self.assertEqual(report['total_rows'], 3)
        self.assertEqual(self.chart(report, 'branch')['values'], ['3'])
        self.assertEqual(report['shown_rows_limit'], 1)

    def test_timing_query_growth_is_batched_and_does_not_query_holidays_per_case(self):
        def queries():
            with CaptureQueriesContext(connection) as captured:
                self.run_report()
            return len(captured)
        self.case()
        baseline = queries()
        for _ in range(12):
            record = self.case()
            self.event(record, 'application_imported', self.now - timedelta(days=1))
        self.assertLessEqual(queries(), baseline + 1)

    def test_invalid_filters_and_selections_fail_safely(self):
        self.case()
        for f in [{'date_mode': 'month', 'month': '2026-99'}, {'date_mode': 'custom', 'from': '2026-02-30', 'to': '2026-03-01'},
                  {'chart_key': 'raw_sql'}, {'chart_key': 'stages', 'bucket_key': 'jbl_visit', 'series_key': 'Not allowed'},
                  {'dimension_field': 'national_id'}]:
            with self.subTest(filters=f), self.assertRaises(PortalReportingError):
                self.run_report(filters=f)

    def test_export_keeps_untrusted_names_as_text(self):
        record = self.case(customer_name='=1+1')
        workbook = load_workbook(BytesIO(export_curated_report(preset='pipeline', filters={}, user=self.user, access=None)))
        self.assertEqual(workbook['Data']['B2'].value, record.customer_name)
        self.assertEqual(workbook['Data']['B2'].data_type, 's')
