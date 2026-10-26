"""Real native-session operations; synthetic database, all external IO denied."""
import json
from datetime import timedelta
from unittest.mock import patch

from django.test import TestCase, override_settings
from django.utils import timezone

from core.models import AccessGrant, CaseUpdate, ComplaintCaseControl, ComplaintCaseEvent, ComplianceAuditEvent, OperationalLocation, BranchServiceArea, UserMiniAppPreference, ParsedMessage
from core import tests_browser_complaints as reads
from staff_auth.tests import TEST_KEY


@override_settings(ROOT_URLCONF='core.test_browser_complaint_urls', STAFF_BROWSER_AUTH_ENABLED=True,
                   STAFF_BROWSER_COMPLAINTS_ENABLED=True, STAFF_MFA_ENCRYPTION_KEYS=[TEST_KEY],
                   EMAIL_BACKEND='django.core.mail.backends.locmem.EmailBackend', SECURE_SSL_REDIRECT=False)
class BrowserComplaintOperationsTests(TestCase):
    setUpTestData = classmethod(reads.BrowserComplaintReadTests.setUpTestData.__func__)
    setUp = reads.BrowserComplaintReadTests.setUp
    post_auth = reads.BrowserComplaintReadTests.post_auth
    login = reads.BrowserComplaintReadTests.login

    @property
    def base(self):
        return f'/api/staff/complaints/operations/v1/groups/{self.group.pk}/'

    def post(self, path, data, csrf=True):
        return self.client.post(self.base + path, json.dumps(data), content_type='application/json',
            **({'HTTP_X_CSRFTOKEN': self.client.cookies['csrftoken'].value} if csrf else {}))

    def grant_role(self, role):
        AccessGrant.objects.create(user=self.user, workflow='complaint_cases', role=role, group_configuration=self.group)

    def test_session_csrf_and_feature_switch_are_independent(self):
        self.assertEqual(self.client.get(self.base + 'bootstrap/').status_code, 401)
        self.login()
        self.assertEqual(self.post('cases/create/', {}, csrf=False).status_code, 403)
        with self.settings(STAFF_BROWSER_COMPLAINTS_ENABLED=False):
            self.assertEqual(self.client.get(self.base + 'bootstrap/').status_code, 404)
            # Restricted v1 remains unaffected by the operations kill switch.
            self.assertEqual(self.client.get(f'/api/staff/complaints/v1/groups/{self.group.pk}/cases/').status_code, 200)

    def test_preferences_confirm_exact_retry_and_reject_stale_overwrite(self):
        self.login()
        original=self.client.get(self.base+'preferences/').json()
        desired={**original,'default_filters':{'status':'resolved', 'branch':''}}
        body={'preferences':desired,'expected':original}
        saved=self.post('preferences/',body)
        self.assertEqual(saved.status_code,200,saved.content)
        audits = ComplianceAuditEvent.objects.filter(action='miniapp.preference.updated').count()
        retry = self.post('preferences/',body)
        self.assertEqual(retry.status_code, 200, retry.content)
        self.assertEqual(retry.json(),saved.json())
        self.assertEqual(ComplianceAuditEvent.objects.filter(action='miniapp.preference.updated').count(), audits)
        rejected=self.post('preferences/',{'preferences':{**desired,'default_filters':{'status':'pending'}},'expected':original})
        self.assertEqual(rejected.status_code,409)
        self.assertEqual(self.client.get(self.base+'preferences/').json(),saved.json())

    def test_full_search_projection_and_detail_are_explicit(self):
        self.login()
        data = self.post('cases/', {'q': '999999999999'}).json()
        self.assertEqual(len(data['items']), 7)
        item = data['items'][0]
        self.assertEqual(item['customer_id'], '999999999999')
        self.assertNotIn('case_id', item)
        self.assertNotIn(self.cases[0].message_id, json.dumps(data))
        response = self.client.get(self.base + f'cases/{self.cases[0].pk}/')
        self.assertEqual(response.status_code, 200, response.content)
        self.assertIn('history', response.json()['case'])
        self.assertNotIn('EXCLUDED_RAW_MESSAGE', response.content.decode())

    def test_queue_order_and_navigation_span_pages_without_crossing_scope(self):
        from core.services.complaint_cases import ensure_case_control
        self.login()
        started = timezone.now() - timedelta(days=100)
        for index in range(24):
            case = reads.synthetic_case(self.group, 2000 + index, timestamp=started + timedelta(days=index))
            ensure_case_control(case)
        oldest = self.post('cases/', {'order': 'oldest', 'page': 1}).json()
        selected = oldest['items'][-1]['key']
        response = self.post('cases/', {'order': 'oldest', 'selected': selected})
        self.assertEqual(response.status_code, 200, response.content)
        navigation = response.json()['navigation']
        self.assertEqual(navigation['position'], 10)
        self.assertEqual(navigation['next']['page'], 2)
        self.assertEqual(navigation['previous']['page'], 1)
        newest = self.post('cases/', {'order': 'newest'}).json()
        self.assertNotEqual(oldest['items'][0]['key'], newest['items'][0]['key'])
        foreign = self.post('cases/', {'order': 'oldest', 'selected': str(self.other_case.pk)}).json()
        self.assertIsNone(foreign['navigation'])
        self.assertEqual(self.post('cases/', {'order': 'unsafe-column'}).status_code, 400)

    def test_recovery_totals_are_exact_beyond_the_visible_list_and_group_scoped(self):
        from core.services.complaint_cases import ensure_case_control
        self.login()
        for index in range(55):
            case = reads.synthetic_case(self.group, 1000 + index)
            control = ensure_case_control(case)
            control.sync_status = 'failed' if index < 3 else 'suspended' if index < 5 else 'pending'
            control.save(update_fields=['sync_status'])
        foreign = ensure_case_control(self.other_case)
        foreign.sync_status = 'failed'
        foreign.save(update_fields=['sync_status'])
        response = self.client.get(self.base + 'recovery/')
        self.assertEqual(response.status_code, 200)
        data = response.json()
        self.assertEqual(data['counts'], {'total': 55, 'pending': 50, 'failed': 3, 'suspended': 2})
        self.assertEqual(len(data['issues']), 50)
        self.assertEqual([row['status'] for row in data['issues'][:3]], ['failed'] * 3)
        self.assertNotIn(str(self.other_case.pk), {row['key'] for row in data['issues']})
        self.assertEqual(self.client.get(self.base.replace(str(self.group.pk), str(self.other.pk)) + 'recovery/').status_code, 403)

    def test_recovery_binding_survives_relogin_but_not_account_security_change(self):
        self.login()
        first = self.client.get(self.base + 'bootstrap/').json()['continuity_key']
        self.assertRegex(first, r'^[a-f0-9]{64}$')
        self.assertEqual(self.client.get(self.base + 'bootstrap/').json()['continuity_key'], first)
        from allauth.mfa.recovery_codes.internal.auth import RecoveryCodes
        from staff_auth.tests import PASSWORD
        codes = RecoveryCodes.activate(self.user).get_unused_codes()
        def sign_in(code):
            self.post_auth('logout', {})
            self.assertEqual(self.post_auth('login', {'username': self.user.username, 'password': PASSWORD}).status_code, 401)
            self.assertEqual(self.post_auth('mfa', {'code': code}).status_code, 200)
        sign_in(codes[0])
        self.assertEqual(self.client.get(self.base + 'bootstrap/').json()['continuity_key'], first)
        from staff_auth.models import StaffBrowserAdmission
        from staff_auth.lifecycle import apply_change
        from django.contrib.auth import get_user_model
        row = StaffBrowserAdmission.objects.get(user=self.user)
        apply_change(actor=get_user_model().objects.get(username='synthetic_admission_root'),
            target_user_id=self.user.pk, action='revoke_sessions', expected_revision=row.revision,
            request_key='synthetic-recovery-revoke', reason='Synthetic recovery boundary test')
        self.assertEqual(self.client.get(self.base + 'bootstrap/').status_code, 401)
        sign_in(codes[1])
        self.assertNotEqual(self.client.get(self.base + 'bootstrap/').json()['continuity_key'], first)

    def test_foreign_group_and_telegram_credentials_fail_closed(self):
        self.login()
        response = self.client.get(self.base + f'cases/{self.other_case.pk}/')
        self.assertEqual(response.status_code, 400)
        self.assertEqual(self.client.get(self.base + 'bootstrap/', HTTP_X_TELEGRAM_INIT_DATA='synthetic').status_code, 403)
        self.grant.delete()
        AccessGrant.objects.create(user=self.user, workflow='jawabu_portal', role='IT', group_configuration=self.group)
        self.assertEqual(self.client.get(self.base + 'bootstrap/').status_code, 403)

    def test_create_field_errors_preserve_no_partial_case(self):
        self.login()
        before = ComplaintCaseEvent.objects.count()
        result = self.post('cases/create/', {'client_request_id': 'synthetic-create-001'})
        self.assertEqual(result.status_code, 400, result.content)
        self.assertIn('customer_phone', result.json()['field_errors'])
        self.assertEqual(ComplaintCaseEvent.objects.count(), before)

    def test_valid_create_replays_and_uses_canonical_locations(self):
        county = OperationalLocation.objects.create(location_type='county', code='SYNTHETIC-COUNTY', name='Synthetic County')
        constituency = OperationalLocation.objects.create(location_type='sub_county', code='SYNTHETIC-SUB', name='Synthetic Constituency', parent=county)
        BranchServiceArea.objects.create(branch=self.branch, area=county)
        self.grant.branch = self.branch.name
        self.grant.branch_ref = self.branch
        self.grant.save(update_fields=['branch', 'branch_ref'])
        from core.models import ComplaintCategory
        ComplaintCategory.objects.create(key='synthetic-browser-service', label='Synthetic Browser Service')
        self.login()
        category = self.client.get(self.base + 'bootstrap/').json()['categories'][0]
        body = {'client_request_id': 'synthetic-valid-create-001', 'client_name': 'Synthetic Browser Customer',
            'customer_phone': '0700000098', 'customer_id': '90000098', 'branch_region': self.branch.code,
            'county': county.code, 'sub_county': constituency.code, 'village': 'Synthetic landmark',
            'complaint_category': category, 'complaint_description': 'Synthetic inspection required.'}
        first = self.post('cases/create/', body)
        self.assertEqual(first.status_code, 201, first.content)
        retry = self.post('cases/create/', body)
        self.assertEqual(retry.status_code, 200, retry.content)
        self.assertEqual(first.json()['case']['key'], retry.json()['case']['key'])
        self.assertEqual(first.json()['case']['county'], county.name)
        self.assertEqual(self.post('cases/create/', {**body, 'client_name': 'Changed request'}).status_code, 400)

    def test_comment_retries_revision_conflicts_and_canonical_audit(self):
        self.grant_role('HB_STAFF')
        self.login()
        detail = self.client.get(self.base + f'cases/{self.cases[0].pk}/').json()['case']
        path = f'cases/{self.cases[0].pk}/comment/'
        body = {'client_request_id': 'synthetic-comment-001', 'expected_revision': detail['revision'],
                'comment_text': 'Synthetic inspection scheduled.'}
        first = self.post(path, body)
        self.assertEqual(first.status_code, 200, first.content)
        self.assertEqual(self.post(path, body).status_code, 200)
        self.assertEqual(CaseUpdate.objects.filter(client_request_id=body['client_request_id']).count(), 1)
        self.assertEqual(self.post(path, {**body, 'comment_text': 'Different'}).status_code, 400)
        self.assertEqual(self.post(path, {**body, 'client_request_id': 'synthetic-comment-002'}).status_code, 409)
        event = ComplianceAuditEvent.objects.filter(action='complaint.case.commented').latest('occurred_at')
        self.assertEqual(event.metadata['request_context']['channel'], 'browser')

    def test_role_denied_and_publication_failure_does_not_rollback_resolution(self):
        self.login()
        path = f'cases/{self.cases[0].pk}/resolve/'
        body = {'client_request_id': 'synthetic-resolve-001', 'expected_revision': 1,
                'resolution_text': 'Synthetic repair completed.'}
        self.assertEqual(self.post(path, body).status_code, 403)
        self.grant_role('HB_STAFF')
        detail = self.client.get(self.base + f'cases/{self.cases[0].pk}/').json()['case']
        body['expected_revision'] = detail['revision']
        self.group.complaint_sheet_projection_enabled = True
        self.group.save(update_fields=['complaint_sheet_projection_enabled'])
        with patch('core.services.complaint_cases.update_sheet_case', side_effect=RuntimeError('synthetic outage')):
            response = self.post(path, body)
        self.assertEqual(response.status_code, 200, response.content)
        self.cases[0].refresh_from_db()
        self.assertEqual(self.cases[0].complaint_status, 'Closed')
        self.assertEqual(response.json()['case']['sync_status'], 'failed')

    def test_server_draft_revision_and_no_attachment_storage(self):
        self.login()
        body = {'fields': {'client_name': 'Synthetic draft'}, 'revision': 0}
        self.assertEqual(self.post('draft/', body).status_code, 200)
        self.assertEqual(self.post('draft/', body).status_code, 409)
        draft = self.client.get(self.base + 'draft/').json()['draft']
        self.assertEqual(draft['fields'], body['fields'])
        self.assertEqual(self.post('draft/', {'action': 'discard', 'revision': 0}).status_code, 409)
        self.assertEqual(self.post('draft/', {'action': 'discard', 'key': draft['key'], 'revision': draft['revision']}).status_code, 200)
        self.assertIsNone(self.client.get(self.base + 'draft/').json()['draft'])
        self.assertEqual(self.post('draft/', {'action': 'discard', 'key': draft['key'], 'revision': draft['revision']}).status_code, 200)
        self.post('draft/', body)
        # The new row starts at revision 1 too; a stale discard must not delete it.
        self.assertEqual(self.post('draft/', {'action': 'discard', 'key': draft['key'], 'revision': draft['revision']}).status_code, 409)

    def test_reports_and_export_require_owning_workflow_permission(self):
        self.login()
        self.assertEqual(self.post('reports/', {}).status_code, 403)
        self.grant_role('IT')
        response = self.post('reports/', {})
        self.assertEqual(response.status_code, 200, response.content)
        self.assertIn('by_time', response.json()['summary'])
        with patch('core.services.complaint_register.export_register_xlsx', return_value=(b'synthetic-xlsx', 7)):
            result = self.post('reports/export/', {'scope': 'filtered', 'filters': {}, 'client_request_id': 'synthetic-export-001'})
        self.assertEqual(result.status_code, 200, result.content)
        self.assertIn('filename=', result['Content-Disposition'])

    def pdf_evidence(self, case=None):
        from core.models import ComplaintCaseEvidence
        case = case or self.cases[0]
        update = CaseUpdate.objects.create(parsed_message=case, group_id=case.group_id,
            old_status='Open', new_status='Open', updated_by='Synthetic officer')
        return ComplaintCaseEvidence.objects.create(parsed_message=case, case_update=update,
            group_id=case.group_id, mime_type='application/pdf', original_filename='synthetic.pdf',
            drive_file_id='synthetic-private-file', upload_status='success')

    def test_evidence_raster_preview_is_authenticated_csrf_protected_and_scoped(self):
        item = self.pdf_evidence()
        self.assertEqual(self.post(f'evidence/{item.pk}/preview/', {}, csrf=False).status_code, 403)
        self.login()
        self.grant_role('IT')
        self.assertEqual(self.post(f'evidence/{item.pk}/preview/', {}, csrf=False).status_code, 403)
        other = self.pdf_evidence(self.other_case)
        with patch('core.services.order_approval.GoogleDriveMediaStorage.download') as download:
            self.assertEqual(self.post(f'evidence/{other.pk}/preview/', {}).status_code, 400)
            for page in (0, -1, True, '1', 10001):
                self.assertEqual(self.post(f'evidence/{item.pk}/preview/', {'page': page}).status_code, 400)
            download.assert_not_called()

    def test_evidence_preview_renders_requested_page_without_browser_pdf_plugin(self):
        from io import BytesIO
        from reportlab.pdfgen.canvas import Canvas
        self.login()
        self.grant_role('IT')
        item = self.pdf_evidence()
        source = BytesIO()
        canvas = Canvas(source)
        for label in ('Synthetic first page', 'Synthetic second page'):
            canvas.drawString(50, 700, label)
            canvas.showPage()
        canvas.save()
        with self.settings(GOOGLE_DRIVE_MEDIA_FOLDER_ID='synthetic-private-folder'), patch('core.services.order_approval.GoogleDriveMediaStorage.download', return_value=source.getvalue()):
            response = self.post(f'evidence/{item.pk}/preview/', {'page': 2})
        self.assertEqual(response.status_code, 200, response.content)
        self.assertEqual(response.json()['page'], 2)
        self.assertEqual(response.json()['page_count'], 2)
        self.assertTrue(response.json()['image'].startswith('data:image/jpeg;base64,/9j/'))
        self.assertEqual(response['Cache-Control'], 'no-store, private')
        self.assertNotIn('synthetic-private-file', response.content.decode())

    def test_evidence_preview_failure_has_a_download_recovery_message(self):
        self.login()
        self.grant_role('IT')
        item = self.pdf_evidence()
        with patch('core.services.order_approval.GoogleDriveMediaStorage.download', side_effect=TimeoutError):
            response = self.post(f'evidence/{item.pk}/preview/', {})
        self.assertEqual(response.status_code, 503)
        self.assertIn('download the original', response.json()['error'])

    def test_evidence_projection_and_bounded_image_thumbnail_are_private(self):
        from io import BytesIO
        from PIL import Image
        self.login()
        self.grant_role('IT')
        item = self.pdf_evidence()
        item.mime_type = 'image/png'
        item.size = 1234
        item.uploaded_by = 'Synthetic uploader'
        item.save(update_fields=['mime_type', 'size', 'uploaded_by'])
        result = self.client.get(self.base + f'cases/{self.cases[0].pk}/').json()['case']['evidence'][0]
        self.assertEqual(result['size'], 1234)
        self.assertEqual(result['uploaded_by'], 'Synthetic uploader')
        self.assertNotIn('drive_file_id', result)
        source = BytesIO()
        Image.new('RGB', (1800, 1200), 'navy').save(source, 'PNG')
        with self.settings(GOOGLE_DRIVE_MEDIA_FOLDER_ID='synthetic-private-folder'), patch('core.services.order_approval.GoogleDriveMediaStorage.download', return_value=source.getvalue()):
            response = self.post(f'evidence/{item.pk}/preview/', {'thumbnail': True})
        self.assertEqual(response.status_code, 200, response.content)
        import base64
        preview = Image.open(BytesIO(base64.b64decode(response.json()['image'].split(',')[1])))
        self.assertLessEqual(max(preview.size), 320)
        self.assertEqual(response.json()['page_count'], 1)
        self.assertNotIn('synthetic-private-file', response.content.decode())

    def test_unsupported_legacy_evidence_is_download_only_and_still_scoped(self):
        self.login()
        self.grant_role('IT')
        item = self.pdf_evidence()
        item.mime_type = 'text/csv'
        item.original_filename = 'synthetic-legacy.csv'
        item.save(update_fields=['mime_type', 'original_filename'])
        with self.settings(GOOGLE_DRIVE_MEDIA_FOLDER_ID='synthetic-private-folder'), patch(
            'core.services.order_approval.GoogleDriveMediaStorage.download', return_value=b'synthetic,value\n'
        ) as source:
            self.assertEqual(self.post(f'evidence/{item.pk}/preview/', {}).status_code, 400)
            source.assert_not_called()
            response = self.client.get(self.base+f'evidence/{item.pk}/?download=1')
            self.assertEqual(response.status_code, 200, response.content)
            self.assertEqual(response['Content-Type'], 'application/octet-stream')
            self.assertIn('attachment;', response['Content-Disposition'])
            self.assertIn('synthetic-legacy.csv', response['Content-Disposition'])
            other = self.pdf_evidence(self.other_case)
            self.other_case.message_id = self.cases[0].message_id
            self.other_case.save(update_fields=['message_id'])
            source.reset_mock()
            self.assertEqual(self.client.get(self.base+f'evidence/{other.pk}/?download=1').status_code, 400)
            self.assertEqual(self.post(f'evidence/{other.pk}/preview/', {}).status_code, 400)
            source.assert_not_called()
            from core.services import complaint_cases
            actor = complaint_cases.staff_actor_for_user(self.group, self.user)
            before = ComplaintCaseEvent.objects.count()
            with self.assertRaises(complaint_cases.ComplaintCaseError):
                complaint_cases.record_evidence_preview(self.group, actor, other)
            self.assertEqual(ComplaintCaseEvent.objects.count(), before)

    def test_saved_preferences_are_in_bootstrap_and_workflow_isolated(self):
        self.login()
        before = UserMiniAppPreference.objects.count()
        bootstrap = self.client.get(self.base + 'bootstrap/').json()
        self.assertEqual(bootstrap['preferences']['default_filters'], {})
        self.assertEqual(UserMiniAppPreference.objects.count(), before)
        response = self.post('preferences/', {'preferences': {'default_filters': {'status': 'resolved', 'branch': 'Synthetic Other Branch'}}})
        self.assertEqual(response.status_code, 200, response.content)
        self.assertEqual(self.client.get(self.base + 'bootstrap/').json()['preferences']['default_filters'],
                         {'status': 'resolved', 'branch': 'Synthetic Other Branch'})
        self.assertFalse(UserMiniAppPreference.objects.exclude(workflow='complaint_cases').exists())
        for filters in ({'status': 'payments'}, {'branch': ['invalid']}, {'status': {'x': 'all'}}, {'search': 'not permitted'}):
            self.assertEqual(self.post('preferences/', {'preferences': {'default_filters': filters}}).status_code, 400)

    def test_keyset_history_does_not_duplicate_when_new_activity_arrives(self):
        self.login()
        case = self.cases[0]
        for index in range(61):
            CaseUpdate.objects.create(parsed_message=case, group_id=case.group_id, old_status='Open',
                new_status='Open', source='mini_app_comment', updated_by='Synthetic history author',
                resolution_text=f'Synthetic older note {index:03d}')
        first = self.post(f'cases/{case.pk}/history/', {}).json()
        self.assertEqual(len(first['items']), 25)
        self.assertTrue(first['next_cursor'])
        fresh = CaseUpdate.objects.create(parsed_message=case, group_id=case.group_id, old_status='Open',
            new_status='Open', source='mini_app_comment', updated_by='Synthetic newer author', resolution_text='Synthetic new note')
        rows = list(first['items'])
        cursor = first['next_cursor']
        while cursor:
            result = self.post(f'cases/{case.pk}/history/', {'cursor': cursor})
            self.assertEqual(result.status_code, 200, result.content)
            data = result.json()
            rows.extend(data['items'])
            cursor = data['next_cursor']
        self.assertEqual(len(rows), 62)  # Original fixture comment and the 61 older notes.
        self.assertEqual(len({row['id'] for row in rows}), 62)
        self.assertNotIn(str(fresh.pk), {row['id'] for row in rows})
        self.assertNotIn('EXCLUDED_RAW_UPDATE', json.dumps(rows))
        native = self.client.get(self.base + f'cases/{case.pk}/').json()['case']
        self.assertEqual(len(native['history']), 25)
        self.assertTrue(native['history_truncated'])
        self.assertTrue(native['history_next_cursor'])

    def test_history_cursor_is_case_bound_tamper_evident_and_reauthorizes(self):
        self.login()
        case = self.cases[0]
        for _ in range(26):
            CaseUpdate.objects.create(parsed_message=case, group_id=case.group_id,
                                      old_status='Open', new_status='Open', source='mini_app_comment')
        cursor = self.post(f'cases/{case.pk}/history/', {}).json()['next_cursor']
        for value in (cursor + 'x', 13, ['invalid']):
            self.assertEqual(self.post(f'cases/{case.pk}/history/', {'cursor': value}).status_code, 400)
        self.assertEqual(self.post(f'cases/{self.cases[1].pk}/history/', {'cursor': cursor}).status_code, 400)
        self.grant.delete()
        self.assertEqual(self.post(f'cases/{case.pk}/history/', {'cursor': cursor}).status_code, 403)

    def test_expired_history_can_be_refreshed_and_legacy_note_appears_only_once(self):
        self.login()
        case = self.cases[1]
        case.resolution_details = 'Synthetic historical resolution'
        case.save(update_fields=['resolution_details'])
        for _ in range(51):
            CaseUpdate.objects.create(parsed_message=case, group_id=case.group_id,
                                      old_status='Open', new_status='Open', source='mini_app_comment')
        # Equal timestamps still use a stable unique tie-breaker.
        CaseUpdate.objects.filter(parsed_message=case).update(created_at=timezone.now())
        first = self.post(f'cases/{case.pk}/history/', {}).json()
        self.assertNotIn('Synthetic historical resolution', [row['note'] for row in first['items']])
        with patch('django.core.signing.TimestampSigner.timestamp', return_value='1'):
            expired = self.post(f'cases/{case.pk}/history/', {}).json()['next_cursor']
        response = self.post(f'cases/{case.pk}/history/', {'cursor': expired})
        self.assertEqual(response.status_code, 400, response.content)
        self.assertEqual(response.json()['field_errors']['history_cursor'], 'invalid')
        data = self.post(f'cases/{case.pk}/history/', {}).json()
        rows = list(data['items'])
        while data['next_cursor']:
            data = self.post(f'cases/{case.pk}/history/', {'cursor': data['next_cursor']}).json()
            rows.extend(data['items'])
        self.assertEqual(len({row['id'] for row in rows}), 52)
        self.assertEqual(sum(row['note'] == 'Synthetic historical resolution' for row in rows), 1)

    def test_email_is_delegated_only_after_current_owning_authorization(self):
        from django.http import JsonResponse
        self.login()
        with patch('report_delivery.views.report_export_response', return_value=JsonResponse({'status': 'queued'})) as dispatch:
            self.assertEqual(self.post('reports/email/', {'email': 'synthetic@example.invalid'}).status_code, 403)
            dispatch.assert_not_called()
            self.grant_role('IT')
            self.assertEqual(self.post('reports/email/', {'email': 'synthetic@example.invalid'}).status_code, 200)
            self.assertEqual(dispatch.call_args.args[:2], (self.user, 'complaint_cases'))

    def report_fixture(self):
        """Different report/resolution cohorts; never touch an operational DB."""
        now = timezone.now()
        self.enterContext(patch('django.utils.timezone.now', return_value=now))
        start = now.replace(day=1, hour=0, minute=0, second=0, microsecond=0)
        pairs = [(start - timedelta(days=6), start, 24),
                 (start - timedelta(days=5), start + timedelta(hours=1), 24),
                 (start, start + timedelta(hours=2), 24)]
        for index, (case, (reported, resolved, target)) in enumerate(zip(self.cases, pairs)):
            case.timestamp, case.date_resolved, case.complaint_status = reported, resolved, 'Closed'
            case.save(update_fields=['timestamp', 'date_resolved', 'complaint_status'])
            ComplaintCaseControl.objects.create(parsed_message=case,
                reference_number=f'CMP-TIMING-{index}', sla_started_at=reported, sla_target_hours=target)
        for index, case in enumerate(self.cases[3:], start=3):
            case.timestamp = start + timedelta(hours=3)
            case.save(update_fields=['timestamp'])
            ComplaintCaseControl.objects.create(parsed_message=case,
                reference_number=f'CMP-TIMING-{index}', sla_started_at=case.timestamp)
        return {'date_from': start.date().isoformat(), 'date_to': now.date().isoformat(),
                'date_basis': 'reported', 'branch': self.cases[0].branch_region}

    def test_report_cohorts_and_all_timing_drills_use_canonical_facts(self):
        from core.services.complaint_register import complaint_report_page, complaint_report_summary
        from core.services.complaint_timing import AGE_BUCKETS
        self.login()
        self.grant_role('IT')
        filters = self.report_fixture()
        summary = self.post('reports/', {**filters, 'granularity': 'month'}).json()['summary']
        self.assertEqual(summary, complaint_report_summary(filters=filters, granularity='month'))
        self.assertEqual(sum(row['received'] for row in summary['activity']), 5)
        self.assertEqual(sum(row['resolved'] for row in summary['activity']), 3)
        self.assertEqual(summary['timing']['target_unavailable'], 0)
        self.assertEqual(summary['timing']['resolution_count'], 3)
        self.assertEqual(summary['timing']['median_resolution_hours'], 121)
        metrics = [('received', 'reported', ''), ('closures', 'closures', ''),
                   ('reported_closed', 'reported', ''), ('reported_open', 'reported', ''),
                   ('resolution', 'resolved', ''), ('on_time', 'resolved', ''), ('late', 'resolved', ''),
                   ('hb_response', 'response', ''), ('reopened', 'reopened', '')]
        metrics += [('open_age', 'reported', row[0]) for row in AGE_BUCKETS]
        for metric, basis, value in metrics:
            selection = {'metric': metric, 'date_basis': basis, 'metric_value': value}
            with self.subTest(metric=metric, value=value):
                response = self.post('reports/', {**filters, 'selection': selection})
                self.assertEqual(response.status_code, 200, response.content)
                self.assertEqual(response.json()['summary'], summary)
                actual = response.json()['table']
                for row in actual['results']:
                    key = row.pop('key')
                    if key:
                        self.assertTrue(ParsedMessage.objects.filter(pk=key, group_id=self.group.group_id).exists())
                self.assertEqual(actual, complaint_report_page(filters={**filters, **selection}, page_size=25))
        resolved = self.post('reports/', {**filters, 'date_basis': 'resolved'}).json()
        self.assertEqual(resolved['summary']['total'], 3)
        self.assertEqual(resolved['summary']['activity'], summary['activity'])

    def test_report_sort_comparison_and_links_keep_canonical_scope(self):
        from core.services.complaint_register import complaint_report_page, complaint_report_summary
        self.login()
        self.grant_role('IT')
        filters = self.report_fixture()
        data = self.post('reports/', {**filters, 'sort': 'customer_name'}).json()
        canonical = complaint_report_page(filters=filters, sort='customer_name')
        self.assertEqual([row['customer_name'] for row in data['table']['results']],
                         [row['customer_name'] for row in canonical['results']])
        start = timezone.datetime.fromisoformat(filters['date_from']).date()
        end = timezone.datetime.fromisoformat(filters['date_to']).date()
        previous = complaint_report_summary(filters={**filters,
            'date_from': (start - (end - start + timedelta(days=1))).isoformat(),
            'date_to': (start - timedelta(days=1)).isoformat()})
        self.assertEqual(data['previous'], {key: previous[key] for key in
            ('total', 'pending', 'resolved', 'card_timing', 'card_coverage')})
        self.assertEqual(self.post('reports/', {'sort': 'raw_message'}).status_code, 400)
        self.assertEqual(self.client.get(self.base + f'cases/{self.other_case.pk}/').status_code, 400)
        ComplaintCaseControl.objects.create(parsed_message=self.other_case, reference_number='CMP-FOREIGN-LINK')
        foreign = self.post('reports/', {'branch': self.other_case.branch_region}).json()['table']['results']
        self.assertEqual([row['key'] for row in foreign if row['complaint_id'] == 'CMP-FOREIGN-LINK'], [''])

    def test_report_table_search_does_not_change_summary_or_comparison(self):
        self.login()
        self.grant_role('IT')
        filters = self.report_fixture()
        baseline = self.post('reports/', filters).json()
        wanted = baseline['table']['results'][0]['complaint_id']
        response = self.post('reports/', {**filters, 'table_search': wanted})
        self.assertEqual(response.status_code, 200, response.content)
        data = response.json()
        self.assertEqual(data['summary'], baseline['summary'])
        self.assertEqual(data['previous'], baseline['previous'])
        self.assertEqual([row['complaint_id'] for row in data['table']['results']], [wanted])
        self.assertEqual(data['table']['page_size'], 25)

    def test_report_metrics_expose_cohort_coverage_and_comparable_previous_timings(self):
        self.login()
        self.grant_role('IT')
        data = self.post('reports/', self.report_fixture()).json()
        coverage = data['summary']['card_coverage']
        self.assertEqual(coverage['first_response']['eligible'], data['summary']['total'])
        self.assertEqual(coverage['resolution']['eligible'], data['summary']['resolved'])
        self.assertEqual(coverage['resolution']['contributing'], data['summary']['card_timing']['resolution_count'])
        self.assertIn('card_timing', data['previous'])
        self.assertIn('card_coverage', data['previous'])
        self.assertTrue(all(row['target_outcome'] in {'on_time', 'late', 'no_target', 'not_resolved', 'not_recorded'}
                            for row in data['table']['results']))

    def test_report_selection_validation_and_empty_timing(self):
        self.login()
        self.grant_role('IT')
        for selection in ([], {'metric': ['late']}, {'unexpected': 'value'},
                          {'metric': 'late', 'date_basis': 'reported'},
                          {'metric': 'open_age', 'metric_value': 'invalid'}):
            with self.subTest(selection=selection):
                self.assertEqual(self.post('reports/', {'selection': selection}).status_code, 400)
        result = self.post('reports/', {'branch': 'No such synthetic branch'}).json()
        self.assertEqual(result['table']['count'], 0)
        self.assertIsNone(result['summary']['timing']['on_time_percent'])
        self.assertIsNone(result['summary']['timing']['median_resolution_hours'])
        self.assertEqual(result['summary']['resolution_trend'], [])

    def test_filtered_export_retains_the_exact_chart_selection(self):
        from io import BytesIO
        from openpyxl import load_workbook
        self.login()
        self.grant_role('IT')
        filters = {**self.report_fixture(), 'date_basis': 'resolved', 'metric': 'late'}
        expected = self.post('reports/', {'selection': filters}).json()['table']['results']
        response = self.post('reports/export/', {'scope': 'filtered', 'filters': filters,
                                               'client_request_id': 'browser-timing-excel-001'})
        self.assertEqual(response.status_code, 200, response.content)
        workbook = load_workbook(BytesIO(response.content))
        self.assertEqual(workbook.active.max_row, len(expected) + 1)
        self.assertEqual({row[1].value for row in list(workbook.active.rows)[1:]},
                         {row['complaint_id'] for row in expected})

    def test_voice_and_preferences_reject_invalid_values_without_provider_calls(self):
        self.login()
        with patch('core.services.portal_voice.create_transcription') as provider:
            self.assertEqual(self.post('voice/', {'field_name': []}).status_code, 400)
            self.assertEqual(self.post('voice/', {'field_name': 'complaint_resolution_note'}).status_code, 403)
            provider.assert_not_called()
        self.assertEqual(self.post('preferences/', {'preferences': {'default_screen': 'payments'}}).status_code, 400)
        self.assertEqual(self.client.get(self.base + 'recovery/').status_code, 200)
