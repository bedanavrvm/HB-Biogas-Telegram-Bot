"""Read reviewed values by header, never by cell colour or an old review flag."""
import re
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from io import BytesIO
from zipfile import ZipFile, BadZipFile
from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException
from django.utils import timezone
from core.services.sheet_publication import OPERATIONAL_FIELDS
from core.services.identifiers import normalize_kenyan_phone, validate_kenyan_national_id
from core.services.jawabu_validation import parse_business_date


def token(value):
    return re.sub(r'[^a-z0-9]', '', str(value or '').lower())


def literal(value):
    if isinstance(value, datetime): return value.date().isoformat()
    return value.isoformat() if isinstance(value, date) else str(value).strip() if value is not None else ''


META = {'sourcerow': 'source_row', 'legacyoutcomestagenotinferred': 'legacy_outcome',
        'legacyoutcome': 'legacy_outcome', 'legacyreason': 'legacy_reason', 'legacyreasoncode': 'legacy_reason_code',
        'decisionstage': 'decision_stage', 'reviewernotes': 'reviewer_notes', 'caseid': 'case_id', 'legacyprogressnotes': 'legacy_progress_notes'}
ALIASES = {token(alias): f.name for f in OPERATIONAL_FIELDS['jawabu_master'] for alias in f.aliases}
DATES = {'hbg_visit_date', 'jbl_visit_date', 'requisition_date', 'hb_deposit_date', 'invoice_date', 'installation_date', 'commissioning_date'}
MONEY = {'invoice_amount', 'discount', 'deposit_paid_hbg'}


def read_workbook(content):
    if len(content) > 20 * 1024 * 1024:
        raise ValueError('Workbook exceeds 20 MB.')
    try:
        with ZipFile(BytesIO(content)) as archive:
            if sum(entry.file_size for entry in archive.infolist()) > 256*1024*1024:
                raise ValueError('Expanded workbook exceeds 256 MB.')
        wb = load_workbook(BytesIO(content), read_only=True, data_only=False)
        cached_wb = load_workbook(BytesIO(content), read_only=True, data_only=True)
    except (BadZipFile, InvalidFileException, KeyError) as exc:
        raise ValueError('Select a readable XLSX workbook.') from exc
    result = []; seen = set()
    try:
        for ws in wb:
            if ws.title.startswith('Projected ') or ws.max_row < 2:
                continue
            headers = next(ws.iter_rows(max_row=1, values_only=True)); bindings = {}; used = set()
            for i, header in enumerate(headers):
                label = str(header or '').strip()
                if '| Original' in label:
                    continue
                if '|' in label:
                    parts = [p.strip() for p in label.split('|')]
                    if parts[-1] != 'Cleaned':
                        continue
                    label = parts[-2]
                if token(label) in {'reviewstate', 'redcells', 'reviewsummary'}: continue
                # The review deliberately preserves opaque legacy columns (including HB customer IDs).
                # They must not acquire a different app meaning just because a label resembles an alias.
                legacy_only = bool('|' in str(header) and str(header).split('|')[0].strip() in {'A','C','E','F','N','R','S','T','AA','AI','AR','AS','AT','AU','AV'})
                field = 'legacy_sign_date' if token(label) == 'signdate' else (
                    'legacy:'+label if legacy_only else META.get(token(label)) or ALIASES.get(token(label)))
                if not field and label: field = 'legacy:'+label
                if field:
                    if field in used:
                        raise ValueError(f'{ws.title}: duplicate mapped header {label}.')
                    bindings[i] = field; used.add(field)
            if 'source_row' not in used or 'customer_name' not in used:
                continue
            if ws.max_row > 100000 or max(bindings) > 255:
                raise ValueError('Data sheets must have at most 100,000 physical rows and mapped columns within the first 256 columns.')
            if ws.title not in {'Master Data','Master Data Compare','Eco-conserve','Eco-conserve Compare'}:
                raise ValueError('Retain the Master Data / Eco-conserve source sheet identities before importing.')
            origin = {'Master Data Compare': 'Master Data', 'Eco-conserve Compare': 'Eco-conserve'}.get(ws.title, ws.title)
            for cells, caches in zip(ws.iter_rows(min_row=2, max_col=max(bindings)+1), cached_wb[ws.title].iter_rows(min_row=2, max_col=max(bindings)+1)):
                values = {}; derived_legacy = set()
                for i, field in bindings.items():
                    value = cells[i].value
                    if cells[i].data_type == 'f':
                        value = caches[i].value
                        if (value is None or caches[i].data_type == 'e') and not field.startswith('legacy:'):
                            raise ValueError(f'{origin}: {field} has no usable cached formula result. Recalculate in Excel or supply reviewed values.')
                        values['legacy formula:'+field] = str(cells[i].value)
                        if field.startswith('legacy:'): derived_legacy.add(field)
                    values[field] = literal(value)
                if not any(v not in ('', '-') for k, v in values.items() if k != 'source_row' and not k.startswith('legacy formula:') and k not in derived_legacy):
                    continue
                try:
                    row = int(values.pop('source_row'))
                    if row < 2: raise ValueError()
                except (ValueError, TypeError):
                    raise ValueError(f'{origin}: every record needs its original Source Row (2 or greater).')
                identity = (origin, row)
                if identity in seen: raise ValueError(f'{origin}: repeated Source Row {row}.')
                seen.add(identity); result.append((origin, row, values))
                if len(result) > 50000: raise ValueError('Workbook exceeds 50,000 records.')
    finally:
        wb.close()
        cached_wb.close()
    if not result: raise ValueError('No data sheet with Source Row and Customer Name was found.')
    return result


def propose(values):
    fields = {}; issues = []; evidence = []
    for key, raw in values.items():
        if raw in ('', '-', None): continue
        value = raw
        if key in DATES:
            parsed = parse_business_date(raw)
            if not parsed: issues.append(f'{key}: invalid date'); continue
            value = parsed.isoformat()
        elif key in MONEY:
            try:
                value = Decimal(str(raw).replace(',', ''))
                if not value.is_finite() or value < 0 or value.as_tuple().exponent < -2: raise InvalidOperation()
                value = str(value)
            except (InvalidOperation, ValueError): issues.append(f'{key}: invalid amount'); continue
        elif key in {'primary_phone', 'secondary_phone'}:
            value = normalize_kenyan_phone(raw)
            if not value: issues.append(f'{key}: complete contact needs confirmation'); continue
        elif key == 'national_id':
            value = validate_kenyan_national_id(raw)
            if not value: issues.append('national_id: identity needs confirmation'); continue
        fields[key] = value
    if not fields.get('hbg_visit_date') and values.get('legacy_sign_date'):
        parsed = parse_business_date(values['legacy_sign_date'])
        if parsed:
            fields['hbg_visit_date'] = parsed.isoformat(); evidence.append('Explicit complete Sign Date alias used because HBG Visit Date was absent')
    for ref in ('order_number', 'payment_number'):
        if fields.get(ref) and not re.fullmatch(r'(?:(?:HB|ECO)-|#)?[1-9][0-9]*', fields[ref], re.I):
            issues.append(f'{ref}: reference meaning/allocator needs review; original retained')
            fields.pop(ref)
    outcome = str(values.get('legacy_outcome', '')).upper()
    outcome = {'DEFERRED': 'DEFERRED / ON HOLD', 'ON HOLD': 'DEFERRED / ON HOLD'}.get(outcome, outcome)
    stage = {'jblvisit': 'jbl_visit', 'credit': 'credit', 'creditanalysis': 'credit', 'final': 'final_review', 'finalreview': 'final_review', 'headofruralreview': 'final_review'}.get(token(values.get('decision_stage', '')))
    for field, explicit_stage in [('final_decision', 'final_review'), ('credit_decision', 'credit')]:
        if str(values.get(field, '')).upper() in {'APPROVED', 'REJECTED', 'DEFERRED / ON HOLD'}:
            stage = explicit_stage; outcome = str(values[field]).upper(); evidence.append(f'Explicit {field}')
            break
    text = '\n'.join(str(values.get(k, '')) for k in ('case_comment_history', 'legacy_progress_notes', 'reviewer_notes', 'pending_installation_comment', 'pending_commissioning_comment', 'cs_remarks'))
    state = 'legacy_review'
    # Only explicit event wording can establish a stage from comments.
    if not stage:
        hits = []
        for key, pattern in [('jbl_visit', r'\bJBL visit\b'), ('credit', r'\bcredit (?:analysis|decision)\b'), ('final_review', r'\b(?:final approval|Head of Rural)\b')]:
            for match in re.finditer(pattern + r'[^\n.]{0,25}?\b(approved|rejected|deferred)\b', text, re.I):
                if re.search(r'\b(?:not|never|awaiting|pending|if|once|will|would|may|should|to be)\b', match[0], re.I):
                    continue
                decision = 'DEFERRED / ON HOLD' if match[1].lower() == 'deferred' else match[1].upper()
                hits.append((key, decision))
        hits = list(dict.fromkeys(hits))
        if len(hits) == 1 and (not outcome or outcome == hits[0][1]):
            stage, outcome = hits[0]; evidence.append('Explicit matching stage and decision in Additional Comments/reviewer notes')
        elif hits: issues.append('Comment decisions need reconciliation with the reported outcome/current stage.')
    if stage and outcome in {'APPROVED', 'REJECTED', 'DEFERRED / ON HOLD'} and not (stage == 'jbl_visit' and outcome == 'APPROVED'):
        state = {'APPROVED': {'jbl_visit': 'credit', 'credit': 'final_review', 'final_review': 'order'}[stage],
                 'REJECTED': 'rejected', 'DEFERRED / ON HOLD': 'deferred'}[outcome]
        evidence.append(f'Legacy decision {outcome} at {stage}')
    if outcome in {'OPTED FOR CASH', 'OPTED FOR OTHER PARTNER', 'WITHDRAWN'}:
        state = 'withdrawn'; evidence.append('Explicit reported closure')
    if fields.get('order_number'):
        # A reference establishes an order, not an accepted scan or HB release.
        state = 'ordered'; evidence.append('Recorded legacy order reference; HB handover may need confirmation')
    physical = values.get('installation_status', '').upper()
    commissioned = values.get('commissioning_status', '').lower()
    if physical in {'INSTALLED', 'CLOSED'} or commissioned == 'commissioned':
        state = 'ordered'; evidence.append('Explicit reported physical completion')
    if physical == 'OPEN' and fields.get('order_number'): evidence.append('Explicit installation pending with recorded order')
    chronology_issue = False
    if physical == 'INSTALLED' or commissioned == 'commissioned':
        for field in ('installation_date', 'commissioning_date'):
            if fields.get(field) and date.fromisoformat(fields[field]) > timezone.localdate():
                issues.append(f'{field}: reported completion is in the future'); chronology_issue = True
        if fields.get('installation_date') and fields.get('commissioning_date') and fields['commissioning_date'] < fields['installation_date']:
            issues.append('Commissioning predates installation; confirm the historical dates.'); chronology_issue = True
    if outcome in {'REJECTED', 'DEFERRED / ON HOLD', 'WITHDRAWN'} and (fields.get('order_number') or physical == 'INSTALLED' or commissioned == 'commissioned'):
        state = 'legacy_review'; issues.append('Decision and later fulfilment facts conflict; confirm current position.')
    if chronology_issue: state = 'legacy_review'
    return {'fields': fields, 'state': state, 'stage': stage or '', 'evidence': evidence, 'issues': issues,
            'outcome': outcome, 'hb_known': state != 'legacy_review' and (physical in {'INSTALLED', 'CLOSED'} or commissioned == 'commissioned' or physical == 'OPEN' and bool(fields.get('order_number')))}
