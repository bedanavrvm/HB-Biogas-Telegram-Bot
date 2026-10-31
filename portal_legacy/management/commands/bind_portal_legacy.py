"""Review live row bindings first; apply only that exact, unchanged manifest."""
import hashlib
import json
from pathlib import Path
from django.contrib.auth import get_user_model
from django.core.management.base import BaseCommand, CommandError
from django.db import transaction
from core.services.sheet_columns import NamedColumns
from core.services.sheets import GoogleSheetsService
from core.services.telegram_identity import user_access
from core.services.requisition_partners import fulfillment_partner_for_farmer, PARTNER_ECO
from portal_legacy.models import LegacyBatch, LegacyRow, LegacyEvent
from portal_legacy.services import authorize


def digest(value):
    return hashlib.sha256(json.dumps(value, ensure_ascii=False, separators=(',', ':')).encode()).hexdigest()


def row_digest(values, case_column):
    data = list(values)
    if case_column and len(data) >= case_column: data[case_column-1] = ''
    while data and not str(data[-1]).strip(): data.pop()
    return digest(data)


class Command(BaseCommand):
    help = 'Preview reviewed legacy Case ID bindings; --apply writes only Case IDs and enables queued projection.'

    def add_arguments(self, parser):
        parser.add_argument('--batch', required=True)
        parser.add_argument('--actor', required=True, type=int)
        parser.add_argument('--manifest', required=True)
        parser.add_argument('--apply', action='store_true')
        parser.add_argument('--note', default='')
        parser.add_argument('--limit', type=int, default=50)

    def handle(self, *args, **opts):
        actor = get_user_model().objects.get(pk=opts['actor'], is_active=True)
        batch = LegacyBatch.objects.select_related('group').get(pk=opts['batch'])
        access = user_access(actor, 'jawabu_portal', group_configuration=batch.group)
        authorize(actor, access, batch.group)
        if not 1 <= opts['limit'] <= 50: raise CommandError('Limit must be 1–50.')
        path = Path(opts['manifest']); workflow = batch.group.workflow or {}
        sheet_id = workflow.get('master_sheet_id') or batch.group.sheet_id
        snapshots = {}

        def snapshot(row):
            name = (workflow.get('eco_conserve_sheet_name') or 'Eco-conserve') if fulfillment_partner_for_farmer(row.farmer) == PARTNER_ECO else (workflow.get('master_sheet_name') or 'Master Data')
            if name not in snapshots:
                service = GoogleSheetsService.get_instance(sheet_id=sheet_id, sheet_name=name)
                if not service.is_available(): raise CommandError('Sheet unavailable; no bindings enabled.')
                sheet = service._sheet
                values = sheet.get_all_values()
                if not values: raise CommandError('Target worksheet has no header.')
                header_row = int(workflow.get('master_header_row') or 1)
                headers = values[header_row-1]
                case_col = NamedColumns(headers).column('Case ID', required=False)
                snapshots[name] = {'sheet': sheet, 'values': values, 'headers': headers, 'header_row': header_row, 'case_col': case_col}
            return name, snapshots[name]

        if not opts['apply']:
            if path.exists(): raise CommandError('Manifest exists. Choose a new path; retain prior reviewed evidence.')
            bindings = []
            for row in batch.rows.filter(committed_at__isnull=False, publication_ready=False).select_related('farmer'):
                authorize(actor, access, batch.group, farmer=row.farmer)
                name, data = snapshot(row)
                target = row.source_row
                if target <= data['header_row'] or target > len(data['values']): raise CommandError('Source row has no live counterpart. Reconcile the row mapping first.')
                bindings.append({'row_id': str(row.pk), 'case_id': str(row.farmer_id), 'source_sheet': row.source_sheet,
                    'source_row': row.source_row, 'target_sheet': name, 'target_row': target,
                    'row_hash': row_digest(data['values'][target-1], data['case_col'])})
            manifest = {'batch': str(batch.pk), 'group': str(batch.group_id), 'spreadsheet': str(sheet_id), 'source_hash': batch.source_hash, 'bindings': bindings}
            path.parent.mkdir(parents=True, exist_ok=True); path.write_text(json.dumps(manifest, indent=2), encoding='utf-8')
            self.stdout.write(f'{len(bindings)} proposed bindings exported. Review source/target rows; no Sheet or workflow writes performed.')
            return
        if not opts['note'].strip(): raise CommandError('Explain the reviewed binding approval with --note.')
        manifest = json.loads(path.read_text(encoding='utf-8'))
        if (manifest.get('batch'), manifest.get('group'), manifest.get('source_hash')) != (str(batch.pk), str(batch.group_id), batch.source_hash):
            raise CommandError('Manifest is for different source facts or ownership; preview again.')
        if manifest.get('spreadsheet') != str(sheet_id):
            raise CommandError('The target spreadsheet changed after review; preview again.')
        from portal_legacy.services import require_finance_reconciled
        require_finance_reconciled(batch.group_id)
        approved = []
        seen = set()
        cases = set()
        for item in manifest['bindings']:
            identity = (item['target_sheet'], int(item['target_row']))
            if identity in seen or item['case_id'] in cases:
                raise CommandError('Duplicate target row or Case ID in the manifest; reconcile all bindings before applying any chunk.')
            seen.add(identity); cases.add(item['case_id'])
        seen.clear()
        for item in manifest['bindings']:
            row = batch.rows.select_related('farmer').get(pk=item['row_id'], committed_at__isnull=False)
            authorize(actor, access, batch.group, farmer=row.farmer)
            if str(row.farmer_id) != item['case_id']: raise CommandError('Manifest case identity changed.')
            if row.publication_ready: continue
            name, data = snapshot(row)
            if name != item['target_sheet']: raise CommandError('Canonical partner/target changed; preview again.')
            target = int(item['target_row']); identity = (name, target)
            if identity in seen or target <= data['header_row'] or target > len(data['values']): raise CommandError('Duplicate or invalid live target row.')
            seen.add(identity)
            current = data['values'][target-1]; col = data['case_col']
            existing = current[col-1] if col and len(current) >= col else ''
            if existing and existing != str(row.farmer_id): raise CommandError('Live row already belongs to another Case ID.')
            if col and any(n != target and len(cells) >= col and cells[col-1] == str(row.farmer_id)
                           for n, cells in enumerate(data['values'], 1)):
                raise CommandError('Case ID already occurs on another live row. Reconcile the existing row rather than duplicating its identity.')
            if row_digest(current, col) != item['row_hash']: raise CommandError('Live row changed after preview. No further bindings applied; review a fresh manifest.')
            approved.append((row, name, target, data))
            if len(approved) == opts['limit']: break
        # This command never replaces customer fields, deletes rows or publishes decisions itself.
        from openpyxl.utils import get_column_letter
        for name, data in snapshots.items():
            members = [entry for entry in approved if entry[1] == name]
            if not members: continue
            fresh_before = data['sheet'].get_all_values()
            for row, _, target, _ in members:
                item = next(i for i in manifest['bindings'] if i['row_id'] == str(row.pk))
                if target > len(fresh_before) or row_digest(fresh_before[target-1], data['case_col']) != item['row_hash']:
                    raise CommandError('Live row changed immediately before binding. Re-preview while manual register edits are paused.')
                existing = fresh_before[target-1][data['case_col']-1] if data['case_col'] and len(fresh_before[target-1]) >= data['case_col'] else ''
                if existing and existing != str(row.farmer_id):
                    raise CommandError('Case ID ownership changed immediately before binding; no overwrite performed.')
            col = data['case_col'] or len(data['headers'])+1
            updates = []
            if not data['case_col']:
                if col > data['sheet'].col_count: data['sheet'].add_cols(col-data['sheet'].col_count)
                updates.append({'range': f'{get_column_letter(col)}{data["header_row"]}', 'values': [['Case ID']]})
            updates.extend({'range': f'{get_column_letter(col)}{target}', 'values': [[str(row.farmer_id)]]} for row, _, target, _ in members)
            data['sheet'].batch_update(updates, value_input_option='RAW')
            fresh = data['sheet'].get_all_values()
            for row, _, target, _ in members:
                if len(fresh) < target or len(fresh[target-1]) < col or fresh[target-1][col-1] != str(row.farmer_id):
                    raise CommandError('Case ID readback failed; publication remains held for unconfirmed records.')
                with transaction.atomic():
                    locked = LegacyRow.objects.select_for_update().get(pk=row.pk)
                    locked.publication_ready = True; locked.revision += 1; locked.save(update_fields=['publication_ready', 'revision'])
                    LegacyEvent.objects.get_or_create(row=locked, request_key='binding:'+digest(manifest)[:64],
                        defaults={'actor': actor, 'action': 'publication_bound', 'payload': {'sheet': name, 'row': target, 'note': opts['note'], 'manifest_hash': digest(manifest)}})
                    from core.services.portal_publication import reserve_farmer_publication
                    reserve_farmer_publication(row.farmer, requested_by=actor, request_id='legacy-binding:'+str(row.pk))
        self.stdout.write(f'{len(approved)} exact Case ID bindings read back and enabled. Repeat the same manifest to resume.')
